from __future__ import annotations
from pathlib import Path
from typing import List, Dict, Any, Optional
from PyQt5 import QtCore, QtGui, QtWidgets
import pyqtgraph as pg
import numpy as np
import tempfile
import time
import logging
import traceback

log = logging.getLogger("secsig")
if not log.handlers:
    _h = logging.StreamHandler()
    _h.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    log.addHandler(_h)
    log.setLevel(logging.DEBUG)

from app.ui.left_panel import DotsCanvas, Stick
from app.ui.mode_dialog import ModeDialog
from app.ui.plot import DataPlot
from app.ui.coords_panel import CoordinatesPanel
from app.io.series_provider import make_series_provider
from app.io.project_store import save_project, load_project, is_project_file, PROJECT_EXT
from app.processing import (
    compute_strain_characteristics,
    StrainResult,
    compute_vibrometry,
    compute_full_vibrometry,
    VibrometryResult,
    FullVibrometryResult,
    BandResult,
    SinusoidalResult,
    estimate_fs_from_time,
    detect_mode_switches,
)


class _PreloadWorker(QtCore.QObject):
    """Предзагрузка X и всех Y в фоновом потоке."""
    progress_value = QtCore.pyqtSignal(int)
    progress_label = QtCore.pyqtSignal(str)
    progress_range = QtCore.pyqtSignal(int)
    finished_signal = QtCore.pyqtSignal(bool)  # canceled

    def __init__(self, provider, file_path: Path, series_list: list):
        super().__init__()
        self._provider = provider
        self._file_path = file_path
        self._series_list = series_list
        self._canceled = False

    @QtCore.pyqtSlot()
    def request_cancel(self):
        self._canceled = True
        if hasattr(self._provider, "request_cancel"):
            try:
                self._provider.request_cancel()
            except Exception:
                pass

    @QtCore.pyqtSlot()
    def run(self):
        try:
            self.progress_label.emit(f"{self._file_path.name}: оценка объёма данных...")
            total_rows = 0
            if hasattr(self._provider, "estimate_rows"):
                total_rows = int(self._provider.estimate_rows())
            else:
                try:
                    x_tmp = self._provider.ensure_x_loaded()
                    total_rows = len(x_tmp) if hasattr(x_tmp, "__len__") else 0
                except Exception:
                    log.exception("preload: ensure_x_loaded for row count")
                    total_rows = 0
        except InterruptedError:
            self._canceled = True
            self.finished_signal.emit(True)
            return
        except Exception:
            log.exception("preload: estimate_rows")
            total_rows = 0

        units_total = total_rows * (1 + len(self._series_list))
        if units_total > 0:
            self.progress_range.emit(units_total)
        self.progress_value.emit(0)

        try:
            self.progress_label.emit(f"{self._file_path.name}: Подготовка X...")
            base_offset = 0

            def cb_x(done, total, _label):
                if units_total > 0:
                    self.progress_value.emit(base_offset + int(min(done, total) if total else done))
                    self.progress_label.emit(f"{self._file_path.name}: X {done}/{total} ({int((done / max(1, total)) * 100)}%)")

            self._provider.ensure_x_loaded(progress=cb_x, total_hint=total_rows or None)
        except InterruptedError:
            self._canceled = True
            self.finished_signal.emit(True)
            return
        except Exception:
            log.exception("preload: ensure_x_loaded")
            self.finished_signal.emit(True)
            return

        base_offset = total_rows if total_rows > 0 else 0
        for si, desc in enumerate(self._series_list):
            if self._canceled:
                self.finished_signal.emit(True)
                return
            series_name = desc['name']
            self.progress_label.emit(f"{self._file_path.name}: {series_name}")
            try:
                def cb_y(done, total, _label):
                    if units_total > 0:
                        self.progress_value.emit(base_offset + int(min(done, total) if total else done))
                        self.progress_label.emit(f"{self._file_path.name}: {series_name} {done}/{total} ({int((done / max(1, total)) * 100)}%)")

                self._provider.load_y(series_name, progress=cb_y, total_hint=total_rows or None)
                base_offset += total_rows if total_rows > 0 else 0
            except InterruptedError:
                self._canceled = True
                self.finished_signal.emit(True)
                return
            except Exception:
                log.exception("preload: load_y %s", series_name)
                continue

        if units_total > 0:
            self.progress_value.emit(units_total)
        self.finished_signal.emit(False)


class _RestoreProjectWorker(QtCore.QObject):
    """Восстановление рядов проекта в фоне с прогрессом."""
    progress_value = QtCore.pyqtSignal(int)
    progress_label = QtCore.pyqtSignal(str)
    progress_range = QtCore.pyqtSignal(int)
    one_series_loaded = QtCore.pyqtSignal(str, str, object, object, object)  # path_str, series_name, x_data, y_data, stick
    finished_signal = QtCore.pyqtSignal(bool)

    def __init__(self, plots: list, providers: dict):
        super().__init__()
        self._plots = list(plots)
        self._providers = providers
        self._canceled = False
        self._scale = 1000

    @QtCore.pyqtSlot()
    def request_cancel(self):
        self._canceled = True
        for p in self._providers.values():
            if hasattr(p, "request_cancel"):
                try:
                    p.request_cancel()
                except Exception:
                    pass

    @QtCore.pyqtSlot()
    def run(self):
        n = len(self._plots)
        self.progress_range.emit(n * self._scale)
        self.progress_value.emit(0)
        for i, plot in enumerate(self._plots):
            if self._canceled:
                self.finished_signal.emit(True)
                return
            path_str = plot.get("source_path", "")
            series_name = plot.get("series_name", "")
            stick = plot.get("stick")
            file_path = Path(path_str)
            provider = self._providers.get(file_path)
            if not provider:
                continue
            self.progress_label.emit(f"Загрузка ряда {i + 1}/{n}...")
            self.progress_value.emit(i * self._scale)
            try:
                def cb(done, total, label):
                    if total and total > 0:
                        frac = min(1.0, (done / total))
                        self.progress_value.emit(int((i + frac) * self._scale))
                    if label:
                        self.progress_label.emit(f"Ряд {i + 1}/{n}: {label}")
                x_data = provider.ensure_x_loaded(progress=cb, total_hint=None)
                if self._canceled:
                    self.finished_signal.emit(True)
                    return
                y_data = provider.load_y(series_name, progress=cb, total_hint=None)
                self.progress_value.emit((i + 1) * self._scale)
                self.one_series_loaded.emit(path_str, series_name, x_data, y_data, stick)
            except InterruptedError:
                self._canceled = True
                self.finished_signal.emit(True)
                return
            except Exception:
                continue
        self.finished_signal.emit(False)


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("SecSig")
        self.resize(1400, 700)

        central = QtWidgets.QWidget(); self.setCentralWidget(central)
        layout = QtWidgets.QHBoxLayout(central)
        layout.setContentsMargins(16, 16, 16, 16); layout.setSpacing(16)

        self.left = DotsCanvas(is_left=True)
        self.right = DataPlot()
        self.coords = CoordinatesPanel()

        self._settings = QtCore.QSettings("SecSig", "SecSig")

        # Открытые файлы: path + список рядов
        self._datasets: List[Dict] = []
        self._plotted_keys: set[str] = set()
        self._providers: Dict[Path, object] = {}
        # LRU индексов графиков при ограничении числа рядов в RAM
        self._memory_mru: List[int] = []
        self._plot_tmp_paths: Dict[int, str] = {}
        self._max_memory_series: int = 7
        self._plot_index_to_source: Dict[int, tuple] = {}
        self._project_path: Path | None = None
        self._project_dirty: bool = False
        self._mode_start_times: List[float] = []
        self._mode_end_times: List[float] = []
        # График, по которому строилась шкала режимов (не в сводных расчётах)
        self._mode_plot_idx: int | None = None
        self._mode_series_name: str | None = None

        self._load_user_settings()

        self.left.stickUpdated.connect(self.on_stick_updated)
        self.right.sigCursorMoved.connect(self.coords.update_cursor_position)
        self.right.sigModeTimesChanged.connect(self._on_mode_times_changed)

        plot_container = QtWidgets.QWidget()
        plot_lay = QtWidgets.QVBoxLayout(plot_container)
        plot_lay.setContentsMargins(0, 0, 0, 0)
        plot_lay.setSpacing(0)
        plot_lay.addWidget(self.right, 1)

        self._x_scrollbar = QtWidgets.QScrollBar(QtCore.Qt.Horizontal)
        self._x_scrollbar.setMinimum(0)
        self._x_scrollbar.setMaximum(0)  # обновится при добавлении данных
        self._x_scrollbar_updating = False  # защита от рекурсивных обновлений
        self._x_scrollbar.valueChanged.connect(self._on_x_scrollbar_moved)
        self.right.sigXRangeChanged.connect(self._on_plot_x_range_changed)
        plot_lay.addWidget(self._x_scrollbar)

        layout.addWidget(self.left, 1)
        layout.addWidget(plot_container, 4)
        layout.addWidget(self.coords, 1)

        self._load_mode = getattr(self, "_load_mode", "lazy")  # "lazy" | "preload_all"
        self._init_menu()
        self._init_chart_toolbar()

        self.left.setContextMenuPolicy(QtCore.Qt.CustomContextMenu)
        self.left.customContextMenuRequested.connect(lambda pos: self._show_plot_context_menu(self.left, pos))
        self.right.setContextMenuPolicy(QtCore.Qt.CustomContextMenu)
        self.right.customContextMenuRequested.connect(lambda pos: self._show_plot_context_menu(self.right, pos))

    _PROGRESS_STYLE = """
        QProgressDialog {
            background-color: #fafafa;
            font-size: 10pt;
        }
        QProgressDialog QLabel {
            color: #333;
            min-width: 280px;
        }
        QProgressBar {
            min-height: 12px;
            max-height: 20px;
            border: 1px solid #c0c0c0;
            border-radius: 4px;
            background: #e8e8e8;
            text-align: center;
        }
        QProgressBar::chunk {
            background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #5ba3f8, stop:1 #2d7dd2);
            border-radius: 3px;
        }
        QPushButton {
            min-width: 70px;
        }
    """

    @staticmethod
    def _format_elapsed(seconds: float) -> str:
        """Форматирует прошедшее время в вид ММ:СС."""
        s = max(0, int(seconds))
        m, s = divmod(s, 60)
        if m > 99:
            m = 99
        return f"{m:02d}:{s:02d}"

    def _create_progress_dialog(
        self,
        title: str,
        label_text: str,
        maximum: int = 0,
        min_duration_ms: int = 0,
    ) -> QtWidgets.QProgressDialog:
        """Прогресс-диалог; после «Отмена» обёртки блокируют обновления со стороны потока (Qt)."""
        dlg = QtWidgets.QProgressDialog(label_text, "Отмена", 0, maximum, self)
        dlg._base_title = title  # type: ignore[attr-defined]
        dlg._user_canceled = False  # type: ignore[attr-defined]
        dlg.setWindowTitle(title)
        dlg.setWindowModality(QtCore.Qt.ApplicationModal)
        dlg.setMinimumDuration(min_duration_ms)
        dlg.setStyleSheet(self._PROGRESS_STYLE)
        dlg.setAutoClose(False)
        dlg.setAutoReset(False)

        _orig_setValue = dlg.setValue
        _orig_setLabelText = dlg.setLabelText
        _orig_setRange = dlg.setRange

        def _on_user_cancel():
            dlg._user_canceled = True  # type: ignore[attr-defined]
            dlg.hide()

        def _safe_setValue(val):
            if not dlg._user_canceled:  # type: ignore[attr-defined]
                _orig_setValue(val)

        def _safe_setLabelText(txt):
            if not dlg._user_canceled:  # type: ignore[attr-defined]
                _orig_setLabelText(txt)

        def _safe_setRange(lo, hi):
            if not dlg._user_canceled:  # type: ignore[attr-defined]
                _orig_setRange(lo, hi)

        dlg.canceled.connect(_on_user_cancel)
        dlg.setValue = _safe_setValue  # type: ignore[assignment]
        dlg.setLabelText = _safe_setLabelText  # type: ignore[assignment]
        dlg.setRange = _safe_setRange  # type: ignore[assignment]

        # Добавляем простой счётчик прошедшего времени в заголовок.
        dlg._start_time = time.monotonic()  # type: ignore[attr-defined]
        timer = QtCore.QTimer(dlg)
        timer.setInterval(1000)
        

        def _update_title():
            try:
                if dlg._user_canceled:  # type: ignore[attr-defined]
                    timer.stop()
                    return
                start = getattr(dlg, "_start_time", None)
                base = getattr(dlg, "_base_title", title)
                if start is None:
                    return
                elapsed = time.monotonic() - float(start)
                dlg.setWindowTitle(f"{base} — {MainWindow._format_elapsed(elapsed)}")
            except Exception:
                timer.stop()

        timer.timeout.connect(_update_title)
        timer.start()
        dlg.finished.connect(timer.stop)
        return dlg

    def _init_menu(self):
        menu = self.menuBar()
        file_menu = menu.addMenu("Файл")
        open_action = file_menu.addAction("Открыть")
        open_action.triggered.connect(self.load_data)
        save_action = file_menu.addAction("Сохранить")
        save_action.setShortcut(QtGui.QKeySequence.Save)  # Ctrl+S
        save_action.triggered.connect(self._save_project)
        save_as_action = file_menu.addAction("Сохранить как...")
        save_as_action.setShortcut(QtGui.QKeySequence("Ctrl+Shift+S"))
        save_as_action.triggered.connect(self._save_project_as)
        export_action = file_menu.addAction("Экспорт")
        export_action.triggered.connect(self._export_results)

        self._recent_menu = file_menu.addMenu("Недавние")
        file_menu.aboutToShow.connect(self._update_recent_menu)

        calc_menu = menu.addMenu("Расчеты")
        for name in [
            "Режим (выгрузить и отобразить)",
            "Калькулятор",
            "Тензометрирование (прочность)",
            "Виброметрирование (вибрации)",
        ]:
            action = calc_menu.addAction(name)
            action.triggered.connect(lambda _checked=False, n=name: self._open_calculation_dialog(n))

        settings_menu = menu.addMenu("Настройки")
        load_mode_action = settings_menu.addAction("Режим загрузки")
        load_mode_action.setToolTip("Настроить, как загружать большие файлы: по одному графику или весь файл сразу")
        load_mode_action.triggered.connect(self._show_load_mode_dialog)

    def _init_chart_toolbar(self):
        """Панель с кнопками управления графиками справа (рядом с Файл/Расчеты)."""
        toolbar = QtWidgets.QToolBar(self)
        toolbar.setMovable(False)
        spacer = QtWidgets.QWidget()
        spacer.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Preferred)
        toolbar.addWidget(spacer)
        btn_zoom_in = QtWidgets.QPushButton("+")
        btn_zoom_in.setToolTip("Приблизить (уменьшить диапазон по X)")
        btn_zoom_in.setFixedSize(32, 24)
        btn_zoom_out = QtWidgets.QPushButton("−")
        btn_zoom_out.setToolTip("Отдалить (увеличить диапазон по X)")
        btn_zoom_out.setFixedSize(32, 24)
        btn_reset = QtWidgets.QPushButton("Сброс")
        btn_reset.setToolTip("Сбросить вид: левая граница 0, правая — максимум по всем графикам")
        btn_reset.setFixedHeight(24)
        toolbar.addWidget(btn_zoom_in)
        toolbar.addWidget(btn_zoom_out)
        toolbar.addWidget(btn_reset)
        btn_zoom_in.clicked.connect(self._zoom_in_charts)
        btn_zoom_out.clicked.connect(self._zoom_out_charts)
        btn_reset.clicked.connect(self._reset_charts_view)
        self.addToolBar(toolbar)

    def _show_plot_context_menu(self, widget: QtWidgets.QWidget, pos: QtCore.QPoint):
        menu = QtWidgets.QMenu(self)
        add_action = menu.addAction("Добавить график")
        add_action.triggered.connect(self.add_plot_dialog)
        clear_action = menu.addAction("Очистить")
        clear_action.triggered.connect(self.clear_plots_only)
        if not self._has_available_series():
            add_action.setEnabled(False)
        menu.exec_(widget.mapToGlobal(pos))

    def _has_available_series(self) -> bool:
        for entry in self._datasets:
            for desc in entry['series']:
                if desc['key'] not in self._plotted_keys:
                    return True
        return False

    def _get_active_plots(self, exclude_mode: bool = False,
                          name_only: bool = False) -> list[tuple[int, str]]:
        plots = []
        if hasattr(self.right, "_plots"):
            for order, idx in enumerate(sorted(self.right._plots.keys())):
                if exclude_mode:
                    if self._mode_plot_idx is not None and idx == self._mode_plot_idx:
                        continue
                    # Дополнительная проверка по имени ряда (индекс может измениться после перезагрузки).
                    if self._mode_series_name:
                        s = self._plot_index_to_source.get(idx)
                        if s and s[1] == self._mode_series_name:
                            continue
                src = self._plot_index_to_source.get(idx)
                if src:
                    path, series_name = src
                    if name_only:
                        label = series_name
                    else:
                        try:
                            fname = Path(path).name
                        except Exception:
                            fname = str(path)
                        label = f"{fname} — {series_name}"
                else:
                    label = f"График {order + 1}"
                plots.append((idx, label))
        return plots

    def get_mode_intervals(self) -> list[tuple[str, float, float]]:
        """Возвращает список активных режимных интервалов: [(label, t_start, t_end), ...]."""
        starts = sorted(self._mode_start_times)
        ends = sorted(self._mode_end_times)
        intervals: list[tuple[str, float, float]] = []
        # Паруем: каждый start_times[i] с ближайшим end_times[j] > start_times[i].
        used_ends: set[int] = set()
        for s in starts:
            for j, e in enumerate(ends):
                if j not in used_ends and e > s:
                    intervals.append((f"Р-{len(intervals) + 1}", s, e))
                    used_ends.add(j)
                    break
        return intervals

    def get_plot_series_data(self, plot_idx: int):
        """Возвращает (x_data, y_data) для графика с индексом plot_idx или None."""
        if not hasattr(self.right, "_plots"):
            return None
        plot_data = self.right._plots.get(plot_idx)
        if not plot_data:
            return None
        x_raw = plot_data.get("x_data")
        y_raw = plot_data.get("y_data")
        if x_raw is None or y_raw is None:
            return None
        x = np.asarray(x_raw, dtype=float)
        y = np.asarray(y_raw, dtype=float)
        if y.size == 0:
            return None
        return (x, y)

    def _build_mode_items(self) -> list[dict]:
        """
        Строит список всех рядов для режима: как уже отрисованных, так и ещё не добавленных.
        Каждый элемент: {label, plot_idx (или None), file, series_name}.
        """
        items: list[dict] = []
        # Обратная карта: (path, series_name) -> plot_idx
        reverse: dict[tuple[Path, str], int] = {}
        for idx, src in self._plot_index_to_source.items():
            path, series_name = src
            try:
                p = Path(path)
            except Exception:
                continue
            reverse[(p.resolve(), series_name)] = idx

        for entry in self._datasets:
            file_path: Path = entry["path"]
            for desc in entry.get("series", []):
                name = desc.get("name")
                if not name:
                    continue
                key = (file_path.resolve(), name)
                plot_idx = reverse.get(key)
                label = f"{file_path.name} — {name}"
                items.append(
                    {
                        "label": label,
                        "plot_idx": plot_idx,
                        "file": str(file_path),
                        "series_name": name,
                    }
                )
        return items

    def _open_calculation_dialog(self, title: str):
        if title == "Тензометрирование (прочность)":
            plots = self._get_active_plots(exclude_mode=True)
            modes = self.get_mode_intervals()
            dlg = TensometryDialog(plots, modes, self)
            dlg.exec_()
        elif title == "Виброметрирование (вибрации)":
            plots = self._get_active_plots(exclude_mode=True)
            modes = self.get_mode_intervals()
            dlg = VibrometryDialog(plots, modes, self)
            dlg.exec_()
        elif title.startswith("Режим"):
            # Упрощённый сценарий: выбор одного ряда и сразу отрисовка режимного графика.
            items = self._build_mode_items()
            if not items:
                QtWidgets.QMessageBox.information(
                    self,
                    "Нет данных",
                    "Нет доступных рядов для анализа режима. Сначала загрузите данные.",
                )
                return
            dlg = ModeDialog(items, self)
            dlg.exec_()
        else:
            # Прочие расчёты используют универсальный диалог.
            items = [
                {"label": label, "plot_idx": idx, "file": None, "series_name": None}
                for idx, label in self._get_active_plots()
            ]
            dlg = CalculationDialog(title, items, self)
            dlg.exec_()

    def get_project_state(self) -> Dict[str, Any]:
        """Состояние проекта для сохранения: источники, графики (путь + ряд + стик), вид по X."""
        sources = []
        seen = set()
        for entry in self._datasets:
            p = entry["path"]
            path_str = str(p.resolve())
            if path_str not in seen:
                seen.add(path_str)
                sources.append({"path": path_str})

        plots = []
        for idx in sorted(getattr(self.right, "_plots", {}).keys()):
            src = self._plot_index_to_source.get(idx)
            if not src:
                continue
            path, series_name = src
            stick = self.left.get_stick(idx)
            if not stick:
                continue
            color_hex = stick.color.name() if stick.color else "#2d5ac8"
            if not color_hex.startswith("#"):
                color_hex = "#2d5ac8"
            plots.append({
                "source_path": str(Path(path).resolve()),
                "series_name": series_name,
                "stick": {
                    "x": stick.x,
                    "y1": stick.y1,
                    "y2": stick.y2,
                    "color_hex": color_hex,
                    "data_min": stick.data_min,
                    "data_max": stick.data_max,
                },
            })

        view = {}
        if hasattr(self.right, "_x_min") and hasattr(self.right, "_x_max"):
            view = {"x_min": float(self.right._x_min), "x_max": float(self.right._x_max)}

        # Режимные маркеры (если пользователь запускал анализ режима).
        mode_markers = {}
        if self._mode_start_times or self._mode_end_times:
            mode_markers = {
                "start_times": [float(t) for t in self._mode_start_times],
                "end_times": [float(t) for t in self._mode_end_times],
                "plot_idx": self._mode_plot_idx,
                "series_name": self._mode_series_name,
            }

        return {"version": 1, "sources": sources, "plots": plots, "view": view,
                "mode_markers": mode_markers}

    def apply_project_state(self, state: Dict[str, Any]) -> None:
        """Восстанавливает workspace из сохранённого состояния. Данные подгружаются из файлов."""
        self.clear_all()
        sources = state.get("sources", [])
        plots = state.get("plots", [])
        if not sources and not plots:
            return

        # Регистрируем источники (провайдеры + список рядов в _datasets)
        for s in sources:
            path_str = s.get("path", "")
            if not path_str:
                continue
            file_path = Path(path_str)
            if not file_path.exists():
                QtWidgets.QMessageBox.warning(
                    self, "Файл не найден",
                    f"Файл источника не найден:\n{file_path}\nРяды из него не будут восстановлены."
                )
                continue
            try:
                provider = make_series_provider(file_path)
                info = provider.list_series()
                self._providers[file_path] = provider
                series_list = [
                    {"key": f"{file_path}|{name}", "file": file_path, "name": name}
                    for name in info.y_names
                ]
                self._datasets.append({"path": file_path, "series": series_list, "x_name": info.x_name})
            except Exception as e:
                QtWidgets.QMessageBox.warning(
                    self, "Ошибка загрузки источника",
                    f"Не удалось открыть файл:\n{file_path}\n\n{e}"
                )

        if not plots:
            view = state.get("view", {})
            if view:
                x_min = view.get("x_min")
                x_max = view.get("x_max")
                if x_min is not None and x_max is not None and hasattr(self.right, "set_x_range_direct"):
                    self.right.set_x_range_direct(float(x_min), float(x_max))
            return

        # Восстанавливаем графики в фоновом потоке с прогрессом (без зависания на больших файлах)
        progress_dlg = self._create_progress_dialog("Открытие проекта", "Восстановление проекта...", len(plots) * 1000, 0)
        worker = _RestoreProjectWorker(plots, self._providers)
        thread = QtCore.QThread(self)
        worker.moveToThread(thread)
        event_loop = QtCore.QEventLoop()

        def on_one_series(path_str: str, series_name: str, x_data, y_data, stick_override):
            file_path = Path(path_str)
            desc = None
            for entry in self._datasets:
                if entry["path"] == file_path:
                    for s in entry["series"]:
                        if s["name"] == series_name:
                            desc = s
                            break
                    break
            if desc is not None:
                try:
                    self._finalize_added_series(desc, x_data, y_data, stick_override=stick_override)
                except Exception as e:
                    QtWidgets.QMessageBox.warning(
                        self, "Ошибка",
                        f"Не удалось добавить ряд «{series_name}» из {file_path.name}\n\n{e}"
                    )

        def on_finished(_canceled: bool):
            thread.quit()
            # Отключаем canceled ПЕРЕД close(), чтобы Qt не эмитил ложную отмену
            try:
                progress_dlg.canceled.disconnect(on_cancel)
            except (TypeError, RuntimeError):
                pass
            progress_dlg.close()
            event_loop.quit()

        def on_cancel():
            worker.request_cancel()
            progress_dlg.close()
            event_loop.quit()

        worker.progress_value.connect(progress_dlg.setValue)
        worker.progress_label.connect(progress_dlg.setLabelText)
        worker.progress_range.connect(lambda m: progress_dlg.setRange(0, max(1, m)))
        worker.one_series_loaded.connect(on_one_series)
        worker.finished_signal.connect(on_finished)
        progress_dlg.canceled.connect(on_cancel)

        thread.started.connect(worker.run)
        thread.finished.connect(thread.deleteLater)
        thread.start()
        event_loop.exec_()

        view = state.get("view", {})
        if view and hasattr(self.right, "set_x_range_direct"):
            x_min = view.get("x_min")
            x_max = view.get("x_max")
            if x_min is not None and x_max is not None:
                self.right.set_x_range_direct(float(x_min), float(x_max))
            elif getattr(self.right, "_plots", {}):
                self.right.set_x_zero_to_data_max(padding_ratio=0.02)

        # Восстанавливаем режимные маркеры (перетаскиваемые палки).
        mode_markers = state.get("mode_markers", {})
        if mode_markers:
            starts = mode_markers.get("start_times", [])
            ends = mode_markers.get("end_times", [])
            if starts or ends:
                self._mode_start_times = [float(t) for t in starts]
                self._mode_end_times = [float(t) for t in ends]
                mpidx = mode_markers.get("plot_idx")
                if mpidx is not None:
                    self._mode_plot_idx = int(mpidx)
                msname = mode_markers.get("series_name")
                if msname:
                    self._mode_series_name = str(msname)
                try:
                    self.right.add_mode_markers(self._mode_start_times, self._mode_end_times)
                except Exception:
                    pass

    def _save_project(self) -> bool:
        """Сохраняет проект в текущий файл (перезапись) или открывает «Сохранить как», если путь не задан."""
        if self._project_path is not None:
            state = self.get_project_state()
            try:
                save_project(state, self._project_path)
                self.setWindowTitle(f"SecSig — {self._project_path.name}")
                self._project_dirty = False
                self._add_to_recent(self._project_path)
                return True
            except Exception as e:
                QtWidgets.QMessageBox.warning(self, "Ошибка сохранения", str(e))
                return False
        return self._save_project_as()

    def _save_project_as(self) -> bool:
        """Сохранить как... Возвращает True, если файл сохранён. Расширение .secsig подставляется автоматически."""
        dlg = QtWidgets.QFileDialog(self, "Сохранить проект", "", "Все файлы (*.*);;Проект SecSig (*.secsig)")
        dlg.setAcceptMode(QtWidgets.QFileDialog.AcceptSave)
        dlg.setDefaultSuffix("secsig")
        if dlg.exec_() != QtWidgets.QDialog.Accepted:
            return False
        paths = dlg.selectedFiles()
        if not paths:
            return False
        path = Path(paths[0])
        if path.suffix.lower() != PROJECT_EXT:
            path = path.with_suffix(PROJECT_EXT)
        state = self.get_project_state()
        try:
            save_project(state, path)
            self._project_path = path
            self.setWindowTitle(f"SecSig — {path.name}")
            self._project_dirty = False
            self._add_to_recent(path)
            return True
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Ошибка сохранения", str(e))
            return False

    def _export_results(self):
        QtWidgets.QMessageBox.information(self, "Экспорт", "Экспорт пока не реализован.")

    _MAX_RECENT_FILES = 10
    _RECENT_FILES_KEY = "recent_files"

    def _get_recent_files(self) -> List[Path]:
        """Список недавно открытых файлов из настроек."""
        try:
            raw = self._settings.value(self._RECENT_FILES_KEY, "")
            if isinstance(raw, str) and raw:
                return [Path(p.strip()) for p in raw.split("\n") if p.strip()]
        except Exception:
            pass
        return []

    def _add_to_recent(self, path: Path) -> None:
        """Добавить путь в начало списка недавних (без дубликатов, не более MAX)."""
        path = path.resolve()
        if not path.exists():
            return
        paths = self._get_recent_files()
        paths = [path] + [p for p in paths if str(p.resolve()) != str(path)]
        paths = paths[: self._MAX_RECENT_FILES]
        self._settings.setValue(self._RECENT_FILES_KEY, "\n".join(str(p) for p in paths))

    def _update_recent_menu(self) -> None:
        """Заполняет подменю «Недавние» перед показом."""
        self._recent_menu.clear()
        recent = self._get_recent_files()
        for p in recent:
            if not p.exists():
                continue
            name = p.name
            if len(str(p)) > 60:
                name = "..." + p.name
            action = self._recent_menu.addAction(name)
            action.setData(str(p))
            action.setToolTip(str(p))
            action.triggered.connect(lambda checked=False, fp=str(p): self._open_recent_file(Path(fp)))
        if recent:
            self._recent_menu.addSeparator()
            clear_action = self._recent_menu.addAction("Очистить список")
            clear_action.triggered.connect(self._clear_recent_list)
        elif not recent:
            a = self._recent_menu.addAction("(нет недавних файлов)")
            a.setEnabled(False)

    def _open_recent_file(self, path: Path) -> None:
        """Открыть файл из списка недавних (проект или данные)."""
        path = Path(path)
        if not path.exists():
            self._remove_from_recent(path)
            QtWidgets.QMessageBox.warning(self, "Файл не найден", f"Файл не найден:\n{path}")
            return
        if is_project_file(path):
            try:
                state = load_project(path)
                self.apply_project_state(state)
                self._project_path = path
                self._project_dirty = False
                self.setWindowTitle(f"SecSig — {path.name}")
                self._add_to_recent(path)
            except Exception as e:
                QtWidgets.QMessageBox.warning(self, "Ошибка открытия проекта", str(e))
        else:
            self._add_to_recent(path)
            self._project_path = None
            self._load_data_from_paths([path])

    def _remove_from_recent(self, path: Path) -> None:
        path = path.resolve()
        paths = [p for p in self._get_recent_files() if str(Path(p).resolve()) != str(path)]
        self._settings.setValue(self._RECENT_FILES_KEY, "\n".join(str(p) for p in paths))

    def _clear_recent_list(self) -> None:
        self._settings.remove(self._RECENT_FILES_KEY)

    def clear_plots_only(self):
        self.left.clear_all()
        self.right.clear_all()
        self.coords.clear_all()
        self._plot_index_to_source.clear()
        for p in list(self._plot_tmp_paths.values()):
            try:
                if Path(p).exists():
                    Path(p).unlink()
            except Exception:
                pass
        self._plot_tmp_paths.clear()
        self._plotted_keys.clear()
        self._memory_mru.clear()
        self._project_dirty = True

    # UI actions
    def load_data(self):
        filepaths, _ = QtWidgets.QFileDialog.getOpenFileNames(
            self,
            "Выберите файлы или проект",
            "",
            "Все файлы (*.*);;Проект SecSig (*.secsig);;Данные (*.xlsx *.csv *.parquet *.parq);;Excel (*.xlsx);;CSV (*.csv);;Parquet (*.parquet *.parq)",
        )
        if not filepaths:
            return

        # Один файл .secsig — открыть проект
        if len(filepaths) == 1 and is_project_file(Path(filepaths[0])):
            path = Path(filepaths[0])
            try:
                state = load_project(path)
                self.apply_project_state(state)
                self._project_path = path
                self._project_dirty = False
                self.setWindowTitle(f"SecSig — {path.name}")
                self._add_to_recent(path)
            except Exception as e:
                QtWidgets.QMessageBox.warning(self, "Ошибка открытия проекта", str(e))
            return

        self._add_to_recent(Path(filepaths[0]))
        self._load_data_from_paths([Path(f) for f in filepaths])

    def _load_data_from_paths(self, filepaths: List[Path]) -> None:
        """Импорт данных по списку путей (без диалога)."""
        log.info("=== _load_data_from_paths START ===")
        log.info(f"  filepaths={[str(p) for p in filepaths]}")
        log.info(f"  _load_mode={self._load_mode!r}")
        if not filepaths:
            log.info("  Пустой список путей — выход")
            return
        self._project_path = None
        imported = 0
        was_canceled = False
        new_file_paths: List[Path] = []
        use_headers_progress = self._load_mode != "preload_all"
        log.info(f"  use_headers_progress={use_headers_progress}")
        progress = None
        if use_headers_progress:
            progress = self._create_progress_dialog("Импорт файлов", "Чтение заголовков...", len(filepaths), 0)
            log.info(f"  Создан прогресс-диалог: maximum={len(filepaths)}")
        for i, file_path in enumerate(filepaths):
            log.info(f"  --- Файл [{i}]: {file_path.name} ---")
            try:
                if progress is not None:
                    progress.setLabelText(f"Чтение: {file_path.name}")
                    progress.setValue(i)
                    QtWidgets.QApplication.processEvents()
                    was_canceled_check = progress.wasCanceled()
                    log.info(f"  progress.wasCanceled() [в цикле] = {was_canceled_check}")
                    if was_canceled_check:
                        log.warning("  >>> BREAK: progress.wasCanceled() == True в цикле!")
                        break
                log.info(f"  Создаю провайдер для {file_path.name}...")
                provider = make_series_provider(file_path)
                log.info(f"  Провайдер: {type(provider).__name__}")
                log.info(f"  Читаю list_series()...")
                info = provider.list_series()
                log.info(f"  x_name={info.x_name!r}, y_names={info.y_names}")
                self._providers[file_path] = provider

                series_list = []
                for name in info.y_names:
                    key = f"{file_path}|{name}"
                    series_list.append({
                        'key': key,
                        'file': file_path,
                        'name': name,
                    })
                self._datasets.append({'path': file_path, 'series': series_list, 'x_name': info.x_name})
                new_file_paths.append(file_path)
                imported += len(info.y_names)
                log.info(f"  imported={imported}, _load_mode={self._load_mode!r}")
                if self._load_mode == "preload_all":
                    log.info(f"  Вызываю _preload_provider_series для {file_path.name}...")
                    canceled = self._preload_provider_series(file_path, provider, series_list, None, 0)
                    log.info(f"  _preload_provider_series вернул canceled={canceled}")
                    if canceled:
                        was_canceled = True
                        log.warning("  >>> BREAK: preload вернул canceled=True!")
                        break
            except Exception as e:
                log.error(f"  EXCEPTION при обработке {file_path.name}: {e}")
                log.error(f"  {traceback.format_exc()}")
                QtWidgets.QMessageBox.warning(self, "Ошибка", f"Не удалось загрузить файл\n{file_path}\n\n{e}")
        if progress is not None:
            progress.setValue(len(filepaths))
            was_canceled_final = progress.wasCanceled()
            log.info(f"  progress.wasCanceled() [после цикла] = {was_canceled_final}")
            if was_canceled_final:
                was_canceled = True
                log.warning("  >>> was_canceled=True из-за progress.wasCanceled() после цикла!")
            progress.close()

        log.info(f"  ИТОГО: was_canceled={was_canceled}, imported={imported}")
        if was_canceled:
            log.warning("  >>> ROLLBACK + показываю 'Отмена'")
            self._rollback_new_imports(new_file_paths)
            QtWidgets.QMessageBox.information(self, "Отмена", "Импорт и предзагрузка отменены пользователем.")
            return

        if imported:
            self._project_dirty = True
            log.info(f"  Успешно! Загружено рядов: {imported}")
            QtWidgets.QMessageBox.information(
                self,
                "Импорт завершён",
                f"Загружено рядов: {imported}. Теперь вы можете добавить их через 'Добавить график…'"
            )
        log.info("=== _load_data_from_paths END ===")

    def add_plot_dialog(self):
        # Подготовка списка доступных, ещё не добавленных рядов
        available = []
        for entry in self._datasets:
            for desc in entry['series']:
                if desc['key'] not in self._plotted_keys:
                    available.append(desc)

        if not available:
            QtWidgets.QMessageBox.information(self, "Нет доступных рядов", "Сначала загрузите файлы, или все ряды уже добавлены.")
            return

        items = [f"{Path(d['file']).name} — {d['name']}" for d in available]
        item, ok = QtWidgets.QInputDialog.getItem(self, "Добавить график", "Выберите ряд:", items, 0, False)
        if not ok:
            return
        chosen_idx = items.index(item)
        chosen = available[chosen_idx]
        self._add_series_to_plot(chosen)

    def _add_series_to_plot(self, desc: Dict):
        file_path: Path = desc['file']
        provider = self._providers.get(file_path)
        if provider is None:
            QtWidgets.QMessageBox.warning(self, "Ошибка", f"Для файла {file_path} не найден провайдер данных.")
            return
        # Если X и Y уже предзагружены — не показываем прогресс, сразу добавляем
        try:
            if getattr(provider, "is_x_cached", lambda: False)() and getattr(provider, "is_y_cached", lambda _n: False)(desc['name']):
                x_data = provider.ensure_x_loaded()
                y_data = provider.load_y(desc['name'])
                self._finalize_added_series(desc, x_data, y_data)
                return
        except Exception:
            pass
        # Запускаем загрузку X/Y в фоновом потоке, чтобы не блокировать GUI
        progress_dlg = self._create_progress_dialog(
            "Загрузка ряда", f"Подготовка «{desc['name']}»...", 0, 0
        )

        class LoadSeriesWorker(QtCore.QObject):
            finished = QtCore.pyqtSignal(object, object, object)  # x, y, err
            setRange = QtCore.pyqtSignal(int, str)
            setValue = QtCore.pyqtSignal(int, str)
            def __init__(self, provider_obj, yname: str):
                super().__init__()
                self.provider_obj = provider_obj
                self.yname = yname
            @QtCore.pyqtSlot()
            def run(self):
                try:
                    # Оценим строки, чтобы показать точный прогресс
                    total_rows = None
                    try:
                        # для CSV есть estimate_rows
                        if hasattr(self.provider_obj, "list_series"):
                            info = self.provider_obj.list_series()
                            xname = getattr(info, "x_name", None)
                        else:
                            xname = None
                        if hasattr(self.provider_obj, "estimate_rows"):
                            total_rows = int(self.provider_obj.estimate_rows(xname))
                    except Exception:
                        total_rows = None
                    if total_rows and total_rows > 0:
                        self.setRange.emit(total_rows * 2, "Подготовка X...")
                    # колбэки прогресса
                    base = 0
                    def progress_x(done, total, _label):
                        try:
                            self.setValue.emit(min(done, total) if total else done, "Подготовка X...")
                        except Exception:
                            pass
                    x_local = self.provider_obj.ensure_x_loaded(progress=progress_x, total_hint=total_rows)
                    base = int(total_rows) if total_rows else len(x_local) if hasattr(x_local, "__len__") else 0
                    def progress_y(done, total, _label):
                        try:
                            self.setValue.emit(base + (min(done, total) if total else done), f"Чтение «{self.yname}»...")
                        except Exception:
                            pass
                    y_local = self.provider_obj.load_y(self.yname, progress=progress_y, total_hint=total_rows)
                    # финальное значение
                    if total_rows and total_rows > 0:
                        self.setValue.emit(total_rows * 2, "Готово")
                    self.finished.emit(x_local, y_local, None)
                except Exception as exc:
                    self.finished.emit(None, None, exc)
            def cancel(self):
                try:
                    # поддерживается CSV-провайдер (стриминговая отмена)
                    if hasattr(self.provider_obj, "request_cancel"):
                        self.provider_obj.request_cancel()
                except Exception:
                    pass

        worker = LoadSeriesWorker(provider, desc['name'])
        thread = QtCore.QThread(self)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)

        _was_canceled = [False]

        def on_finished(x_data, y_data, err):
            # Отключаем canceled ПЕРЕД close(), чтобы Qt не эмитил ложную отмену
            try:
                progress_dlg.canceled.disconnect(on_cancel)
            except (TypeError, RuntimeError):
                pass
            progress_dlg.close()
            thread.quit()
            thread.wait()
            worker.deleteLater()
            thread.deleteLater()
            if _was_canceled[0]:
                return  # пользователь отменил — игнорируем результат
            if err is not None or x_data is None or y_data is None:
                QtWidgets.QMessageBox.warning(self, "Ошибка", f"Не удалось загрузить ряд\n{file_path}\n\n{err}")
                return
            self._finalize_added_series(desc, x_data, y_data)

        def on_cancel():
            _was_canceled[0] = True
            worker.cancel()
            progress_dlg.close()

        progress_dlg.canceled.connect(on_cancel)
        worker.setRange.connect(lambda mx, text: (progress_dlg.setRange(0, mx), progress_dlg.setLabelText(text)))
        worker.setValue.connect(lambda val, text: (progress_dlg.setValue(val), progress_dlg.setLabelText(text)))
        worker.finished.connect(on_finished)
        thread.start()

    def _finalize_added_series(self, desc: Dict, x_data, y_data, stick_override: Optional[Dict[str, Any]] = None):
        file_path: Path = desc['file']
        # Диапазон Y по реальным данным (или из stick_override при восстановлении проекта).
        data_min, data_max = 0.0, 1.0
        if stick_override:
            data_min = float(stick_override.get("data_min", 0.0))
            data_max = float(stick_override.get("data_max", 1.0))
        else:
            try:
                y_arr = np.asarray(y_data)
                n = y_arr.size
                if n > 0:
                    stride = max(1, n // 100000)
                    s = y_arr[::stride]
                    data_min = float(np.nanmin(s))
                    data_max = float(np.nanmax(s))
                    if not np.isfinite(data_min) or not np.isfinite(data_max):
                        data_min, data_max = 0.0, 1.0
                    elif data_min == data_max:
                        data_min, data_max = data_min - 0.5, data_min + 0.5
            except Exception:
                data_min, data_max = 0.0, 1.0

        if stick_override:
            x = int(stick_override.get("x", 50))
            y1 = int(stick_override.get("y1", 50))
            y2 = int(stick_override.get("y2", 350))
            hex_color = stick_override.get("color_hex", "#2d5ac8")
            color = QtGui.QColor(hex_color) if hex_color.startswith("#") else QtGui.QColor(hex_color)
            if not color.isValid():
                color = QtGui.QColor.fromHsv((self.left.stick_count() * 47) % 360, 220, 220)
        else:
            idx = self.left.stick_count()
            hue = (idx * 47) % 360
            color = QtGui.QColor.fromHsv(hue, 220, 220)
            h = self.left.height(); margin = self.left.margin
            stick_h = max(100, h - margin * 2 - 100)
            y1 = int((h - stick_h) / 2); y2 = int(y1 + stick_h)
            avail_w = max(1, self.left.width() - margin * 2 - 60)
            x = int(margin + 30 + (idx * 40) % max(1, avail_w))

        # На шкале слева — только имя ряда, на панели координат — файл + ряд.
        new_stick = Stick(x, y1, y2, color, data_min, data_max, desc['name'])
        idx_added = self.left.add_stick(new_stick)

        full_label = f"{file_path.name} — {desc['name']}"
        self.right.add_or_update_plot(idx_added, x_data, y_data, y1, y2, color, y_min=data_min, y_max=data_max)
        self.coords.update_stick_data(idx_added, color, data_min, data_max, y1, y2, x_data, y_data,
                                      name=full_label)
        self._plotted_keys.add(desc['key'])
        self._plot_index_to_source[idx_added] = (file_path, desc['name'])
        self._project_dirty = True
        try:
            import numpy as _np
            if isinstance(y_data, _np.memmap):
                self._plot_tmp_paths[idx_added] = str(getattr(y_data, 'filename', ''))
        except Exception:
            pass
        self._memory_mru.append(idx_added)
        self._enforce_memory_limit()
        return idx_added

    def on_stick_updated(self, idx: int, y1: int, y2: int):
        self.right.update_plot_v_range(idx, y1, y2)
        if idx in self.coords._sticks_data:
            data = self.coords._sticks_data[idx]
            self.coords.update_stick_data(idx, data['color'], data['data_min'], data['data_max'], y1, y2,
                                          data['x_data'], data['y_data'],
                                          name=data.get('name', ''))

    def clear_all(self):
        self.left.clear_all(); self.right.clear_all(); self.coords.clear_all()
        for p in list(self._plot_tmp_paths.values()):
            try:
                if Path(p).exists():
                    Path(p).unlink()
            except Exception:
                pass
        self._plot_tmp_paths.clear()
        self._plot_index_to_source.clear()
        for prov in list(self._providers.values()):
            try:
                prov.cleanup()
            except Exception:
                pass
        self._memory_mru.clear()
        self._project_dirty = True

    def _zoom_in_charts(self):
        """Приблизить: уменьшить видимый диапазон по X."""
        current_range = self.right._x_half_range
        self.right.set_x_half_range(current_range / 1.5)

    def _zoom_out_charts(self):
        """Отдалить: увеличить видимый диапазон по X."""
        current_range = self.right._x_half_range
        self.right.set_x_half_range(current_range * 1.5)

    def _reset_charts_view(self):
        """Сброс: левая граница 0, правая — максимум среди всех графиков (с небольшим отступом)."""
        self.right.set_x_zero_to_data_max(padding_ratio=0.02)

    def zoom_out_right_panel(self):
        current_range = self.right._x_half_range
        self.right.set_x_half_range(current_range * 1.5)

    def _enforce_memory_limit(self):
        while len(self._memory_mru) > self._max_memory_series:
            idx_to_demote = self._memory_mru.pop(0)
            self._demote_plot_to_memmap(idx_to_demote)

    def _demote_plot_to_memmap(self, idx: int):
        if idx in self._plot_tmp_paths:
            return
        plot_data = self.right._plots.get(idx) if hasattr(self.right, "_plots") else None
        if not plot_data:
            return
        pdialog = self._create_progress_dialog("Оптимизация памяти", "Перенос графика на диск...", 0, 0)
        QtWidgets.QApplication.processEvents()
        x_data = plot_data['x_data']
        y_data = plot_data['y_data']
        y_top = plot_data['y_top']; y_bottom = plot_data['y_bottom']
        color = plot_data['color']
        try:
            try:
                import numpy as _np
                if isinstance(y_data, _np.memmap):
                    self._plot_tmp_paths[idx] = str(getattr(y_data, 'filename', ''))
                    return
            except Exception:
                pass
            tmp = tempfile.NamedTemporaryFile(prefix=f"secsig_y_{idx}_", suffix=".npy", delete=False)
            tmp_path = tmp.name
            tmp.close()
            np.save(tmp_path, np.asarray(y_data))
            y_mem = np.load(tmp_path, mmap_mode="r")
            y_min = float(plot_data.get('y_data_min', 0.0))
            y_max = float(plot_data.get('y_data_max', 1.0))
            self.right.add_or_update_plot(idx, x_data, y_mem, y_top, y_bottom, color, y_min=y_min, y_max=y_max)
            if idx in self.coords._sticks_data:
                data = self.coords._sticks_data[idx]
                self.coords.update_stick_data(idx, data['color'], data['data_min'], data['data_max'],
                                              y_top, y_bottom, x_data, y_mem,
                                              name=data.get('name', ''))
            self._plot_tmp_paths[idx] = tmp_path
        except Exception as e:
            self._memory_mru.append(idx)
        finally:
            pdialog.close()

    def closeEvent(self, event):
        if self._project_dirty:
            msg = QtWidgets.QMessageBox(self)
            msg.setWindowTitle("Подтверждение")
            msg.setText("Проект изменён. Сохранить изменения перед закрытием?")
            btn_save = msg.addButton("Сохранить", QtWidgets.QMessageBox.AcceptRole)
            btn_discard = msg.addButton("Не сохранять", QtWidgets.QMessageBox.DestructiveRole)
            btn_cancel = msg.addButton("Отмена", QtWidgets.QMessageBox.RejectRole)
            msg.setDefaultButton(btn_save)
            msg.exec_()
            clicked = msg.clickedButton()
            if clicked == btn_cancel:
                event.ignore()
                return
            if clicked == btn_save:
                if not self._save_project():
                    event.ignore()
                    return
        try:
            self._save_user_settings()
            self.clear_all()
        finally:
            super().closeEvent(event)

    def _show_load_mode_dialog(self):
        dlg = QtWidgets.QDialog(self)
        dlg.setWindowTitle("Режим загрузки")
        dlg.setMinimumWidth(420)
        lay = QtWidgets.QVBoxLayout(dlg)

        grp = QtWidgets.QGroupBox("Как загружать большие файлы")
        grp_lay = QtWidgets.QVBoxLayout(grp)
        self._rb_lazy_load = QtWidgets.QRadioButton("По одному графику")
        self._rb_lazy_load.setToolTip("Данные читаются только при добавлении ряда на график")
        self._rb_preload = QtWidgets.QRadioButton("Весь файл сразу")
        self._rb_preload.setToolTip("При открытии файла выгружать все ряды в память")
        self._rb_lazy_load.setChecked(self._load_mode == "lazy")
        self._rb_preload.setChecked(self._load_mode == "preload_all")
        grp_lay.addWidget(self._rb_lazy_load)
        grp_lay.addWidget(self._rb_preload)

        row = QtWidgets.QHBoxLayout()
        row.addWidget(grp)
        btn_help = QtWidgets.QPushButton("?")
        btn_help.setFixedSize(28, 28)
        btn_help.setToolTip("Подробнее о режимах")
        btn_help.clicked.connect(lambda: self._show_load_mode_help())
        row.addWidget(btn_help, alignment=QtCore.Qt.AlignTop)
        lay.addLayout(row)

        btns = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        lay.addWidget(btns)

        if dlg.exec_() == QtWidgets.QDialog.Accepted:
            self._load_mode = "preload_all" if self._rb_preload.isChecked() else "lazy"
            self._save_user_settings()
            if self._load_mode == "preload_all":
                ans = QtWidgets.QMessageBox.question(
                    self,
                    "Предзагрузка",
                    "Предзагрузить все серии для уже открытых файлов сейчас?",
                    QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
                    QtWidgets.QMessageBox.No,
                )
                if ans == QtWidgets.QMessageBox.Yes:
                    self._preload_all_series()

    def _show_load_mode_help(self):
        QtWidgets.QMessageBox.information(
            self,
            "Режимы загрузки",
            "<p><b>По одному графику</b></p>"
            "<p>При открытии файла в список попадают только названия рядов. Данные с диска читаются "
            "только когда вы добавляете конкретный ряд на график. Быстрый старт, меньше памяти, "
            "но каждое добавление графика может занять время на больших файлах.</p>"
            "<p><b>Весь файл сразу</b></p>"
            "<p>При открытии файла все ряды выгружаются в память (или на диск через memory-mapping). "
            "Открытие дольше, зато добавление любого графика потом происходит мгновенно. "
            "Удобно, когда нужно часто переключаться между многими рядами одного файла.</p>"
            "<p>Выбор сохраняется и действует при следующем запуске.</p>",
        )

    def _preload_all_series(self):
        has_series = any(len(entry.get('series', [])) > 0 for entry in self._datasets)
        if not has_series:
            QtWidgets.QMessageBox.information(self, "Нет данных", "Сначала загрузите файлы.")
            return
        for entry in self._datasets:
            file_path: Path = entry['path']
            provider = self._providers.get(file_path)
            if not provider:
                continue
            canceled = self._preload_provider_series(file_path, provider, entry.get('series', []), None, 0)
            if canceled:
                QtWidgets.QMessageBox.information(self, "Отмена", f"Предзагрузка файла {Path(file_path).name} отменена.")
                break

    def _preload_provider_series(self, file_path: Path, provider, series_list: list, progress: QtWidgets.QProgressDialog | None = None, current_step: int = 0) -> bool:
        local_progress = progress
        if local_progress is None:
            local_progress = self._create_progress_dialog(
                "Предзагрузка файла", f"{file_path.name}: подготовка...", 0, 0
            )
            local_progress.setValue(0)

        worker = _PreloadWorker(provider, file_path, series_list)
        thread = QtCore.QThread(self)
        worker.moveToThread(thread)
        event_loop = QtCore.QEventLoop()
        _canceled_flag = [False]

        def on_finished(_canceled: bool):
            thread.quit()
            if progress is None:
                try:
                    local_progress.canceled.disconnect(on_cancel)
                except (TypeError, RuntimeError):
                    pass
                local_progress.close()
            event_loop.quit()

        def on_cancel():
            _canceled_flag[0] = True
            worker.request_cancel()
            if progress is None:
                local_progress.close()
            event_loop.quit()

        worker.progress_value.connect(local_progress.setValue)
        worker.progress_label.connect(local_progress.setLabelText)
        worker.progress_range.connect(lambda m: local_progress.setRange(0, max(1, m)))
        worker.finished_signal.connect(on_finished)
        if progress is None:
            local_progress.canceled.connect(on_cancel)

        thread.started.connect(worker.run)
        thread.finished.connect(thread.deleteLater)
        thread.start()
        event_loop.exec_()

        result = _canceled_flag[0] or worker._canceled
        return result

    def _rollback_new_imports(self, file_paths: List[Path]):
        paths_set = set(Path(p) for p in file_paths)
        self._datasets = [d for d in self._datasets if Path(d.get('path')) not in paths_set]
        for p in list(paths_set):
            prov = self._providers.pop(p, None)
            if prov is not None:
                try:
                    prov.cleanup()
                except Exception:
                    pass

    def _ensure_series_plotted_for_mode(self, file_path: Path, series_name: str) -> Optional[int]:
        """
        Гарантирует, что заданный ряд (file_path + series_name) добавлен как график.
        Возвращает индекс графика (plot_idx) или None при ошибке.
        """
        # Если уже отрисован — просто возвращаем существующий индекс.
        for idx, src in self._plot_index_to_source.items():
            path, name = src
            try:
                p = Path(path)
            except Exception:
                continue
            if p.resolve() == file_path.resolve() and name == series_name:
                return idx

        # Ищем описание ряда в _datasets
        desc: Optional[Dict] = None
        for entry in self._datasets:
            if Path(entry.get("path")).resolve() != file_path.resolve():
                continue
            for s in entry.get("series", []):
                if s.get("name") == series_name:
                    desc = s
                    break
            if desc is not None:
                break

        if desc is None:
            QtWidgets.QMessageBox.warning(
                self,
                "Ошибка",
                f"Ряд «{series_name}» для файла {file_path} не найден в загруженных данных.",
            )
            return None

        provider = self._providers.get(file_path)
        if provider is None:
            try:
                provider = make_series_provider(file_path)
                info = provider.list_series()
                self._providers[file_path] = provider
                if not any(Path(d.get("path")).resolve() == file_path.resolve() for d in self._datasets):
                    series_list = []
                    for name in info.y_names:
                        key = f"{file_path}|{name}"
                        series_list.append({"key": key, "file": file_path, "name": name})
                    self._datasets.append({"path": file_path, "series": series_list, "x_name": info.x_name})
            except Exception as e:
                QtWidgets.QMessageBox.warning(self, "Ошибка", f"Не удалось открыть файл\n{file_path}\n\n{e}")
                return None

        try:
            x_data = provider.ensure_x_loaded()
            y_data = provider.load_y(series_name)
        except Exception as e:
            QtWidgets.QMessageBox.warning(
                self,
                "Ошибка",
                f"Не удалось загрузить данные ряда «{series_name}» из файла\n{file_path}\n\n{e}",
            )
            return None

        try:
            idx_added = self._finalize_added_series(desc, x_data, y_data)
        except Exception as e:
            QtWidgets.QMessageBox.warning(
                self,
                "Ошибка",
                f"Не удалось добавить ряд «{series_name}» как график.\n\n{e}",
            )
            return None

        return idx_added

    def _run_mode_analysis_for_item(self, item: Dict[str, Any]) -> None:
        """
        Выполняет анализ режима для выбранного ряда:
        - при необходимости добавляет график,
        - считает моменты переключений,
        - рисует вертикальные линии на основном графике.
        """
        from pathlib import Path as _Path  # локальный импорт, чтобы избежать конфликтов

        if not item:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Не выбран ряд для анализа режима.")
            return

        plot_idx = item.get("plot_idx")
        # Если график ещё не отрисован — добавляем его.
        if plot_idx is None:
            file_str = item.get("file")
            series_name = item.get("series_name")
            if not file_str or not series_name:
                QtWidgets.QMessageBox.warning(self, "Ошибка", "Недостаточно информации о выбранном ряде.")
                return
            file_path = _Path(file_str)
            new_idx = self._ensure_series_plotted_for_mode(file_path, series_name)
            if new_idx is None:
                return
            plot_idx = int(new_idx)
            item["plot_idx"] = plot_idx

        data = self.get_plot_series_data(int(plot_idx))
        if data is None:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Не удалось получить данные выбранного графика.")
            return
        x_data, y_data = data

        try:
            switches, y_min, y_max, mid = detect_mode_switches(x_data, y_data)
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Ошибка", f"Не удалось выполнить анализ режима:\n{e}")
            return

        # классификация переключений: начало/окончание режима
        start_times: list[float] = []
        end_times: list[float] = []
        for sw in switches:
            if sw.from_level < sw.to_level:
                start_times.append(sw.t)
            elif sw.from_level > sw.to_level:
                end_times.append(sw.t)

        # Сохраняем времена режимов, индекс и имя графика-режима, отображаем маркеры.
        self._mode_start_times = list(start_times)
        self._mode_end_times = list(end_times)
        self._mode_plot_idx = int(plot_idx)
        src = self._plot_index_to_source.get(int(plot_idx))
        self._mode_series_name = src[1] if src else None
        if hasattr(self, "right") and hasattr(self.right, "add_mode_markers"):
            try:
                self.right.add_mode_markers(start_times, end_times)
            except Exception:
                pass

    def _on_mode_times_changed(self, start_times: list, end_times: list):
        """Вызывается при перетаскивании режимных линий пользователем."""
        self._mode_start_times = [float(t) for t in start_times]
        self._mode_end_times = [float(t) for t in end_times]
        self._project_dirty = True

    def _on_plot_x_range_changed(self, x_min: float, x_max: float, data_x_max: float):
        """Синхронизирует скролл-бар с текущим видимым диапазоном графика."""
        if self._x_scrollbar_updating:
            return
        self._x_scrollbar_updating = True
        try:
            view_span = x_max - x_min
            total = max(data_x_max, x_max)
            SCALE = 10000
            if total <= 0 or view_span >= total:
                self._x_scrollbar.setMaximum(0)
                self._x_scrollbar.setValue(0)
            else:
                page = int(view_span / total * SCALE)
                sb_max = SCALE - page
                self._x_scrollbar.setPageStep(page)
                self._x_scrollbar.setSingleStep(max(1, page // 10))
                self._x_scrollbar.setMaximum(max(0, sb_max))
                pos = int(x_min / (total - view_span) * sb_max) if (total - view_span) > 0 else 0
                self._x_scrollbar.setValue(max(0, min(sb_max, pos)))
        finally:
            self._x_scrollbar_updating = False

    def _on_x_scrollbar_moved(self, value: int):
        """Перемещает видимый диапазон графика при движении скролл-бара."""
        if self._x_scrollbar_updating:
            return
        self._x_scrollbar_updating = True
        try:
            sb_max = self._x_scrollbar.maximum()
            if sb_max <= 0:
                return
            view_span = self.right._x_max - self.right._x_min
            total = max(self.right._data_x_max, self.right._x_max)
            scroll_range = total - view_span
            if scroll_range <= 0:
                return
            new_x_min = (value / sb_max) * scroll_range
            self.right.set_x_range_direct(new_x_min, new_x_min + view_span)
        finally:
            self._x_scrollbar_updating = False

    def _load_user_settings(self):
        try:
            mode = str(self._settings.value("load_mode", "lazy"))
            self._load_mode = mode if mode in ("lazy", "preload_all") else "lazy"
        except Exception:
            self._load_mode = "lazy"
        try:
            ppp = float(self._settings.value("points_per_pixel", 1.0))
            self.right.set_points_per_pixel(ppp)
        except Exception:
            pass
        try:
            cursor_hz = float(self._settings.value("cursor_hz", 50.0))
            self.right.set_cursor_sample_rate_hz(cursor_hz)
        except Exception:
            pass
        try:
            self._max_memory_series = int(self._settings.value("max_memory_series", self._max_memory_series))
        except Exception:
            pass

    def _save_user_settings(self):
        try:
            self._settings.setValue("load_mode", self._load_mode)
            ppp = float(getattr(self.right, "_max_points_factor", 1.0))
            self._settings.setValue("points_per_pixel", ppp)
            hz = 0.0
            try:
                interval = float(getattr(self.right, "_coords_emit_interval", 0.02))
                hz = (1.0 / interval) if interval > 0 else 50.0
            except Exception:
                hz = 50.0
            self._settings.setValue("cursor_hz", hz)
            self._settings.setValue("max_memory_series", int(self._max_memory_series))
        except Exception:
            pass

class CalculationWorker(QtCore.QObject):
    progress = QtCore.pyqtSignal(int)
    finished = QtCore.pyqtSignal(str)

    def __init__(self, steps: int = 100, delay_s: float = 0.02):
        super().__init__()
        self._steps = int(steps)
        self._delay_s = float(delay_s)

    @QtCore.pyqtSlot()
    def run(self):
        for i in range(self._steps + 1):
            time.sleep(self._delay_s)
            self.progress.emit(i)
        self.finished.emit("Расчет завершен. Логика режима пока не подключена.")


class TensometryBatchWorker(QtCore.QObject):
    """Воркер: для каждой пары (график, режимный интервал) считает StrainResult."""
    progress = QtCore.pyqtSignal(int)
    one_result = QtCore.pyqtSignal(str, str, object)  # graph_label, mode_label, StrainResult
    finished = QtCore.pyqtSignal()

    def __init__(self, items: list, miner_exponent: float = 5.0):
        super().__init__()
        # items: [(graph_label, mode_label, y_slice), ...]
        self._items = list(items)
        self._miner_exponent = float(miner_exponent)

    @QtCore.pyqtSlot()
    def run(self):
        n = len(self._items)
        for i, (graph_label, mode_label, y_arr) in enumerate(self._items):
            y = np.asarray(y_arr, dtype=float)
            r = compute_strain_characteristics(y, miner_exponent=self._miner_exponent)
            self.one_result.emit(graph_label, mode_label, r)
            self.progress.emit(int((i + 1) * 100 / n) if n else 100)
        self.finished.emit()


class VibrometryBatchWorker(QtCore.QObject):
    """Воркер: для каждой пары (график, режимный интервал) считает FullVibrometryResult."""
    progress = QtCore.pyqtSignal(int)
    one_result = QtCore.pyqtSignal(str, str, object)  # graph_label, mode_label, FullVibrometryResult
    finished = QtCore.pyqtSignal()

    def __init__(self, items: list, bands: list[tuple[float, float]] | None = None,
                 base_freqs: list[float] | None = None,
                 delta_f: float = 0.375, af_pct: float = 5.0,
                 ae_pct: float = 5.0, al_pct: float = 50.0,
                 calc_equiv: bool = True, calc_eff: bool = False,
                 segment_length: int | None = None):
        super().__init__()
        self._items = list(items)
        self._bands = bands or []
        self._base_freqs = base_freqs or []
        self._delta_f = delta_f
        self._af_pct = af_pct
        self._ae_pct = ae_pct
        self._al_pct = al_pct
        self._calc_equiv = calc_equiv
        self._calc_eff = calc_eff
        self._segment_length = int(segment_length) if segment_length and segment_length >= 4 else None

    @QtCore.pyqtSlot()
    def run(self):
        n = len(self._items)
        for i, (graph_label, mode_label, x_arr, y_arr) in enumerate(self._items):
            x = np.asarray(x_arr, dtype=float)
            y = np.asarray(y_arr, dtype=float)
            fs = estimate_fs_from_time(x)
            if not np.isfinite(fs) or fs <= 0:
                fs = 1000.0
            r = compute_full_vibrometry(
                y, fs_hz=fs,
                bands=self._bands if self._bands else None,
                base_freqs=self._base_freqs if self._base_freqs else None,
                delta_f=self._delta_f, af_pct=self._af_pct,
                ae_pct=self._ae_pct, al_pct=self._al_pct,
                calc_equiv=self._calc_equiv, calc_eff=self._calc_eff,
                segment_length=self._segment_length,
            )
            self.one_result.emit(graph_label, mode_label, r)
            self.progress.emit(int((i + 1) * 100 / n) if n else 100)
        self.finished.emit()


def _make_checkbox_list(items: list[tuple], empty_msg: str) -> QtWidgets.QListWidget:
    """Создаёт QListWidget с чекбоксами. items: [(data_value, label), ...]."""
    lst = QtWidgets.QListWidget()
    lst.setSelectionMode(QtWidgets.QAbstractItemView.NoSelection)
    for data_val, label in items:
        it = QtWidgets.QListWidgetItem(label)
        it.setFlags(it.flags() | QtCore.Qt.ItemIsUserCheckable)
        it.setCheckState(QtCore.Qt.Unchecked)
        it.setData(QtCore.Qt.UserRole, data_val)
        lst.addItem(it)
    if not items:
        lst.addItem(QtWidgets.QListWidgetItem(empty_msg))
    return lst


def _toggle_all(lst: QtWidgets.QListWidget, state: QtCore.Qt.CheckState):
    for i in range(lst.count()):
        it = lst.item(i)
        if it.data(QtCore.Qt.UserRole) is not None:
            it.setCheckState(state)


def _get_checked(lst: QtWidgets.QListWidget) -> list[tuple]:
    """Возвращает [(data_value, label), ...] для отмеченных элементов."""
    result = []
    for i in range(lst.count()):
        it = lst.item(i)
        dv = it.data(QtCore.Qt.UserRole)
        if dv is not None and it.checkState() == QtCore.Qt.Checked:
            result.append((dv, it.text()))
    return result


def _build_selection_tabs(
    plots: list[tuple[int, str]],
    modes: list[tuple[str, float, float]],
) -> tuple[QtWidgets.QTabWidget, QtWidgets.QListWidget, QtWidgets.QListWidget]:
    """Строит QTabWidget с вкладками «Графики» и «Режимы»."""
    tabs = QtWidgets.QTabWidget()

    graphs_page = QtWidgets.QWidget()
    graphs_lay = QtWidgets.QVBoxLayout(graphs_page)
    btn_row = QtWidgets.QHBoxLayout()
    btn_all = QtWidgets.QPushButton("Выбрать все")
    btn_none = QtWidgets.QPushButton("Снять все")
    btn_row.addWidget(btn_all); btn_row.addWidget(btn_none); btn_row.addStretch(1)
    graphs_lay.addLayout(btn_row)
    graph_list = _make_checkbox_list(
        [(idx, label) for idx, label in plots],
        "Нет доступных графиков (исключая режимный)",
    )
    btn_all.clicked.connect(lambda: _toggle_all(graph_list, QtCore.Qt.Checked))
    btn_none.clicked.connect(lambda: _toggle_all(graph_list, QtCore.Qt.Unchecked))
    graphs_lay.addWidget(graph_list)
    tabs.addTab(graphs_page, "Графики")

    modes_page = QtWidgets.QWidget()
    modes_lay = QtWidgets.QVBoxLayout(modes_page)
    btn_row_m = QtWidgets.QHBoxLayout()
    btn_all_m = QtWidgets.QPushButton("Выбрать все")
    btn_none_m = QtWidgets.QPushButton("Снять все")
    btn_row_m.addWidget(btn_all_m); btn_row_m.addWidget(btn_none_m); btn_row_m.addStretch(1)
    modes_lay.addLayout(btn_row_m)
    mode_items = [(i, f"{label}  [{t0:.4g} – {t1:.4g} с]") for i, (label, t0, t1) in enumerate(modes)]
    mode_list = _make_checkbox_list(mode_items, "Режимы не определены. Сначала выполните анализ режима.")
    btn_all_m.clicked.connect(lambda: _toggle_all(mode_list, QtCore.Qt.Checked))
    btn_none_m.clicked.connect(lambda: _toggle_all(mode_list, QtCore.Qt.Unchecked))
    modes_lay.addWidget(mode_list)
    tabs.addTab(modes_page, "Режимы")

    return tabs, graph_list, mode_list


class TensometryDialog(QtWidgets.QDialog):
    """Тензометрирование: выбор графиков + режимов, параметр Минера, расчёт → экспорт в xlsx."""

    def __init__(self, plots: list[tuple[int, str]],
                 modes: list[tuple[str, float, float]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Тензометрирование (прочность)")
        self.resize(540, 480)
        self._plots = plots
        self._modes = modes
        self._results: list[tuple[str, str, StrainResult]] = []  # (graph, mode, result)
        self._worker_thread = None
        self._worker = None

        layout = QtWidgets.QVBoxLayout(self)

        self._tabs, self._graph_list, self._mode_list = _build_selection_tabs(plots, modes)
        layout.addWidget(self._tabs)

        grp_params = QtWidgets.QGroupBox("Параметры")
        params_lay = QtWidgets.QHBoxLayout(grp_params)
        params_lay.addWidget(QtWidgets.QLabel("Показатель кривой Вёлера (m):"))
        self._miner_spin = QtWidgets.QDoubleSpinBox()
        self._miner_spin.setRange(2.0, 20.0); self._miner_spin.setValue(5.0); self._miner_spin.setDecimals(1)
        params_lay.addWidget(self._miner_spin); params_lay.addStretch(1)
        layout.addWidget(grp_params)

        self._progress = QtWidgets.QProgressBar()
        self._progress.setRange(0, 100); self._progress.setValue(0); self._progress.setVisible(False)
        self._calc_btn = QtWidgets.QPushButton("Рассчитать и сохранить в xlsx")
        self._calc_btn.clicked.connect(self._run_calculation)
        layout.addWidget(self._progress)
        layout.addWidget(self._calc_btn)

    def _run_calculation(self):
        main_win = self.parent()
        if not main_win or not hasattr(main_win, "get_plot_series_data"):
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Нет доступа к данным графика.")
            return

        checked_graphs = _get_checked(self._graph_list)  # [(idx, label), ...]
        checked_modes = _get_checked(self._mode_list)     # [(mode_index, label), ...]
        if not checked_graphs:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Выберите хотя бы один график.")
            return
        if not checked_modes:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Выберите хотя бы один режим.")
            return

        # Собираем пары (graph × mode) и нарезаем данные по режимным интервалам.
        items = []  # [(graph_label, mode_label, y_slice), ...]
        # Для xlsx: маппинг полного имени → короткого (только имя ряда).
        self._short_names: Dict[str, str] = {}
        for g_idx, g_label in checked_graphs:
            src = main_win._plot_index_to_source.get(g_idx)
            self._short_names[g_label] = src[1] if src else g_label
            data = main_win.get_plot_series_data(g_idx)
            if data is None:
                continue
            x_data, y_data = data
            x_arr = np.asarray(x_data)
            y_arr = np.asarray(y_data)
            for m_idx, m_label in checked_modes:
                _label, t0, t1 = self._modes[m_idx]
                i0 = int(np.searchsorted(x_arr, t0, side='left'))
                i1 = int(np.searchsorted(x_arr, t1, side='right'))
                y_slice = y_arr[i0:i1]
                if y_slice.size == 0:
                    continue
                items.append((g_label, _label, y_slice))

        if not items:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Нет данных в выбранных режимных интервалах.")
            return

        self._results.clear()
        self._checked_graphs_labels = [lbl for _, lbl in checked_graphs]
        self._checked_modes_labels = [self._modes[mi][0] for mi, _ in checked_modes]
        self._progress.setVisible(True); self._progress.setValue(0)
        self._calc_btn.setEnabled(False)

        self._worker_thread = QtCore.QThread(self)
        self._worker = TensometryBatchWorker(items, miner_exponent=self._miner_spin.value())
        self._worker.moveToThread(self._worker_thread)
        self._worker_thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._progress.setValue)
        self._worker.one_result.connect(self._collect_result)
        self._worker.finished.connect(self._on_finished)
        self._worker.finished.connect(self._worker_thread.quit)
        self._worker_thread.finished.connect(self._worker_thread.deleteLater)
        self._worker_thread.start()

    def _collect_result(self, graph_label: str, mode_label: str, r: StrainResult):
        self._results.append((graph_label, mode_label, r))

    def _on_finished(self):
        self._progress.setValue(100); self._progress.setVisible(False)
        self._calc_btn.setEnabled(True)
        if not self._results:
            QtWidgets.QMessageBox.information(self, "Готово", "Нет результатов для сохранения.")
            return
        self._export_xlsx()

    def _export_xlsx(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Сохранить результаты", "", "Excel (*.xlsx);;Все файлы (*.*)")
        if not path:
            return

        PARAM_NAMES = ["Мин.", "Макс.", "Среднее",
                       "Экв. полураз.", "Макс. полураз.",
                       "Мин. квазистат.", "Макс. квазистат."]

        wb = Workbook()
        ws = wb.active
        ws.title = "Тензометрирование"
        bold = Font(bold=True)
        center = Alignment(horizontal='center')

        # Группируем результаты: {graph_label: {mode_label: StrainResult}}
        data_map: Dict[str, Dict[str, StrainResult]] = {}
        for g, m, r in self._results:
            data_map.setdefault(g, {})[m] = r

        graphs = self._checked_graphs_labels
        modes = self._checked_modes_labels
        n_params = len(PARAM_NAMES)

        # Заголовок: строка 1 — имена графиков, строка 2 — параметры.
        c = ws.cell(1, 1, "Режим")
        c.font = bold; c.alignment = center
        col = 2
        for g in graphs:
            short = self._short_names.get(g, g)
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + n_params - 1)
            cell = ws.cell(1, col, short)
            cell.font = bold; cell.alignment = center
            for pi, pname in enumerate(PARAM_NAMES):
                c2 = ws.cell(2, col + pi, pname)
                c2.font = bold; c2.alignment = center
            col += n_params

        # Данные: строка за строкой по режимам.
        for ri, mode_label in enumerate(modes):
            row = ri + 3
            mc = ws.cell(row, 1, mode_label)
            mc.font = bold; mc.alignment = center
            col = 2
            for g in graphs:
                r = data_map.get(g, {}).get(mode_label)
                if r:
                    vals = [r.y_min, r.y_max, r.y_mean,
                            r.equivalent_half_range, r.max_half_range,
                            r.min_quasi_static, r.max_quasi_static]
                    for vi, v in enumerate(vals):
                        cell_val = v if not isinstance(v, float) or np.isfinite(v) else ""
                        dc = ws.cell(row, col + vi, cell_val)
                        dc.alignment = center
                col += n_params

        # Автоширина столбцов по содержимому.
        from openpyxl.utils import get_column_letter as _gcl
        for ci in range(1, ws.max_column + 1):
            max_len = 0
            for ri in range(1, ws.max_row + 1):
                val = ws.cell(ri, ci).value
                if val is not None:
                    s = f"{val:.6g}" if isinstance(val, float) else str(val)
                    max_len = max(max_len, len(s))
            ws.column_dimensions[_gcl(ci)].width = max(8, max_len + 3)

        try:
            wb.save(path)
            QtWidgets.QMessageBox.information(self, "Готово", f"Результаты сохранены:\n{path}")
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Ошибка сохранения", str(e))


class _FreqBandTable(QtWidgets.QWidget):
    """Таблица полос частот (f1…f2) с прокруткой и кнопками «Добавить» / «Удалить»."""

    def __init__(self, header_left: str = "f₁", header_right: str = "f₂", parent=None):
        super().__init__(parent)
        lay = QtWidgets.QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)

        self._table = QtWidgets.QTableWidget(0, 2)
        self._table.setHorizontalHeaderLabels([header_left, header_right])
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.verticalHeader().setVisible(True)
        self._table.setMinimumHeight(90)
        lay.addWidget(self._table)

        btn_row = QtWidgets.QHBoxLayout()
        btn_add = QtWidgets.QPushButton("Добавить")
        btn_del = QtWidgets.QPushButton("Удалить")
        btn_add.clicked.connect(self._add_row)
        btn_del.clicked.connect(self._del_row)
        btn_row.addWidget(btn_add)
        btn_row.addWidget(btn_del)
        btn_row.addStretch(1)
        lay.addLayout(btn_row)

    def _add_row(self):
        r = self._table.rowCount()
        self._table.insertRow(r)

    def _del_row(self):
        r = self._table.currentRow()
        if r >= 0:
            self._table.removeRow(r)

    def get_values(self) -> list[tuple[float, float]]:
        result = []
        for r in range(self._table.rowCount()):
            it1 = self._table.item(r, 0)
            it2 = self._table.item(r, 1)
            if it1 and it2:
                try:
                    result.append((float(it1.text().replace(",", ".")),
                                   float(it2.text().replace(",", "."))))
                except ValueError:
                    pass
        return result

    def row_count(self) -> int:
        return self._table.rowCount()

    def has_empty(self) -> bool:
        for r in range(self._table.rowCount()):
            for c in range(2):
                it = self._table.item(r, c)
                if not it or not it.text().strip():
                    return True
        return False


class _SingleFreqTable(QtWidgets.QWidget):
    """Таблица одиночных значений частот с прокруткой и кнопками «Добавить» / «Удалить»."""

    def __init__(self, header: str = "Частота, Гц", parent=None):
        super().__init__(parent)
        lay = QtWidgets.QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)

        self._table = QtWidgets.QTableWidget(0, 1)
        self._table.setHorizontalHeaderLabels([header])
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.verticalHeader().setVisible(True)
        self._table.setMinimumHeight(90)
        lay.addWidget(self._table)

        btn_row = QtWidgets.QHBoxLayout()
        btn_add = QtWidgets.QPushButton("Добавить")
        btn_del = QtWidgets.QPushButton("Удалить")
        btn_add.clicked.connect(self._add_row)
        btn_del.clicked.connect(self._del_row)
        btn_row.addWidget(btn_add)
        btn_row.addWidget(btn_del)
        btn_row.addStretch(1)
        lay.addLayout(btn_row)

    def _add_row(self):
        r = self._table.rowCount()
        self._table.insertRow(r)

    def _del_row(self):
        r = self._table.currentRow()
        if r >= 0:
            self._table.removeRow(r)

    def get_values(self) -> list[float]:
        result = []
        for r in range(self._table.rowCount()):
            it = self._table.item(r, 0)
            if it:
                try:
                    result.append(float(it.text().replace(",", ".")))
                except ValueError:
                    pass
        return result

    def row_count(self) -> int:
        return self._table.rowCount()

    def has_empty(self) -> bool:
        for r in range(self._table.rowCount()):
            it = self._table.item(r, 0)
            if not it or not it.text().strip():
                return True
        return False


class VibrometryDialog(QtWidgets.QDialog):
    """Виброметрирование: С.Ш.В. + синусоидальная вибрация, выбор графиков/режимов → xlsx."""

    def __init__(self, plots: list[tuple[int, str]],
                 modes: list[tuple[str, float, float]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Виброметрирование (вибрации)")
        self.resize(820, 700)
        self._plots = plots
        self._modes = modes
        self._results: list[tuple[str, str, FullVibrometryResult]] = []
        self._worker_thread = None
        self._worker = None

        root = QtWidgets.QVBoxLayout(self)

        self._tabs, self._graph_list, self._mode_list = _build_selection_tabs(plots, modes)
        self._tabs.setMaximumHeight(200)
        root.addWidget(self._tabs)

        params_row = QtWidgets.QHBoxLayout()

        self._shv_cb = QtWidgets.QCheckBox("С. Ш. В.")
        self._shv_cb.setStyleSheet("font-weight: bold; font-size: 13px;")
        self._shv_cb.toggled.connect(self._on_shv_toggled)

        self._shv_frame = QtWidgets.QFrame()
        self._shv_frame.setFrameShape(QtWidgets.QFrame.StyledPanel)
        shv_lay = QtWidgets.QVBoxLayout(self._shv_frame)
        shv_lay.setContentsMargins(6, 4, 6, 4)

        # Подвыбор: СКЗ / Sxx
        cb_row = QtWidgets.QHBoxLayout()
        self._skz_cb = QtWidgets.QCheckBox("СКЗ")
        self._sxx_cb = QtWidgets.QCheckBox("Sxx")
        self._skz_cb.setChecked(True)
        cb_row.addWidget(self._skz_cb)
        cb_row.addWidget(self._sxx_cb)
        cb_row.addStretch(1)
        shv_lay.addLayout(cb_row)

        # Таблица диапазонов частот
        shv_lay.addWidget(QtWidgets.QLabel("Диапазоны частот:"))
        self._bands_table = _FreqBandTable("f₁, Гц", "f₂, Гц")
        shv_lay.addWidget(self._bands_table)

        left_box = QtWidgets.QVBoxLayout()
        left_box.addWidget(self._shv_cb)
        left_box.addWidget(self._shv_frame)
        params_row.addLayout(left_box, 1)

        self._sin_cb = QtWidgets.QCheckBox("Синусоидальная вибрация")
        self._sin_cb.setStyleSheet("font-weight: bold; font-size: 13px;")
        self._sin_cb.toggled.connect(self._on_sin_toggled)

        self._sin_frame = QtWidgets.QFrame()
        self._sin_frame.setFrameShape(QtWidgets.QFrame.StyledPanel)
        sin_lay = QtWidgets.QVBoxLayout(self._sin_frame)
        sin_lay.setContentsMargins(6, 4, 6, 4)

        form = QtWidgets.QFormLayout()
        self._delta_f_spin = QtWidgets.QDoubleSpinBox()
        self._delta_f_spin.setRange(0.001, 1000.0)
        self._delta_f_spin.setValue(0.375)
        self._delta_f_spin.setDecimals(3)
        self._delta_f_spin.setSuffix("")
        form.addRow("Приращение частоты [Δf]:", self._delta_f_spin)

        self._af_spin = QtWidgets.QDoubleSpinBox()
        self._af_spin.setRange(0.01, 100.0)
        self._af_spin.setValue(5.0)
        self._af_spin.setDecimals(2)
        form.addRow("Допуск по частоте Af [%]:", self._af_spin)
        sin_lay.addLayout(form)

        # Эквивалентные амплитуды
        self._equiv_cb = QtWidgets.QCheckBox("Эквивалентные амплитуды")
        sin_lay.addWidget(self._equiv_cb)
        ae_row = QtWidgets.QHBoxLayout()
        ae_row.addWidget(QtWidgets.QLabel("Допуск по энергии Ae [%]:"))
        self._ae_spin = QtWidgets.QDoubleSpinBox()
        self._ae_spin.setRange(0.01, 100.0)
        self._ae_spin.setValue(5.0)
        self._ae_spin.setDecimals(2)
        ae_row.addWidget(self._ae_spin)
        ae_row.addStretch(1)
        sin_lay.addLayout(ae_row)

        # Эффективная амплитуда
        self._eff_cb = QtWidgets.QCheckBox("Эффективная амплитуда")
        sin_lay.addWidget(self._eff_cb)
        al_row = QtWidgets.QHBoxLayout()
        al_row.addWidget(QtWidgets.QLabel("Допуск по уровню Al [%]:"))
        self._al_spin = QtWidgets.QDoubleSpinBox()
        self._al_spin.setRange(0.01, 100.0)
        self._al_spin.setValue(50.0)
        self._al_spin.setDecimals(2)
        al_row.addWidget(self._al_spin)
        al_row.addStretch(1)
        sin_lay.addLayout(al_row)

        # Базовые частоты
        sin_lay.addWidget(QtWidgets.QLabel("Базовые частоты:"))
        self._base_freq_table = _SingleFreqTable("Частота, Гц")
        sin_lay.addWidget(self._base_freq_table)

        right_box = QtWidgets.QVBoxLayout()
        right_box.addWidget(self._sin_cb)
        right_box.addWidget(self._sin_frame)
        params_row.addLayout(right_box, 1)

        root.addLayout(params_row)

        # Начальное состояние: оба блока заблокированы
        self._shv_frame.setEnabled(False)
        self._sin_frame.setEnabled(False)

        self._progress = QtWidgets.QProgressBar()
        self._progress.setRange(0, 100)
        self._progress.setValue(0)
        self._progress.setVisible(False)
        self._calc_btn = QtWidgets.QPushButton("Расчёт")
        self._calc_btn.clicked.connect(self._run_calculation)
        root.addWidget(self._progress)
        root.addWidget(self._calc_btn)

    def _on_shv_toggled(self, checked: bool):
        self._shv_frame.setEnabled(checked)

    def _on_sin_toggled(self, checked: bool):
        self._sin_frame.setEnabled(checked)

    def _validate(self) -> bool:
        if not self._shv_cb.isChecked() and not self._sin_cb.isChecked():
            QtWidgets.QMessageBox.warning(self, "Ошибка",
                                          "Выберите хотя бы один вид анализа (С.Ш.В. или Синусоидальная).")
            return False

        if self._shv_cb.isChecked():
            if not self._skz_cb.isChecked() and not self._sxx_cb.isChecked():
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "С.Ш.В.: выберите хотя бы один параметр (СКЗ или Sxx).")
                return False
            if self._bands_table.row_count() == 0:
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "С.Ш.В.: добавьте хотя бы один диапазон частот.")
                return False
            if self._bands_table.has_empty():
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "С.Ш.В.: заполните все ячейки в таблице диапазонов.")
                return False

        if self._sin_cb.isChecked():
            if not self._equiv_cb.isChecked() and not self._eff_cb.isChecked():
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "Синусоидальная: выберите хотя бы один пункт "
                                              "(Экв. амплитуды или Эфф. амплитуда).")
                return False
            if self._base_freq_table.row_count() == 0:
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "Синусоидальная: добавьте хотя бы одну базовую частоту.")
                return False
            if self._base_freq_table.has_empty():
                QtWidgets.QMessageBox.warning(self, "Ошибка",
                                              "Синусоидальная: заполните все ячейки базовых частот.")
                return False
        return True

    def _run_calculation(self):
        main_win = self.parent()
        if not main_win or not hasattr(main_win, "get_plot_series_data"):
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Нет доступа к данным графика.")
            return

        checked_graphs = _get_checked(self._graph_list)
        checked_modes = _get_checked(self._mode_list)
        if not checked_graphs:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Выберите хотя бы один график.")
            return
        if not checked_modes:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Выберите хотя бы один режим.")
            return
        if not self._validate():
            return

        # Собираем параметры
        bands = self._bands_table.get_values() if self._shv_cb.isChecked() else []
        base_freqs = self._base_freq_table.get_values() if self._sin_cb.isChecked() else []
        calc_equiv = self._equiv_cb.isChecked() if self._sin_cb.isChecked() else False
        calc_eff = self._eff_cb.isChecked() if self._sin_cb.isChecked() else False
        delta_f = self._delta_f_spin.value()
        af_pct = self._af_spin.value()
        ae_pct = self._ae_spin.value()
        al_pct = self._al_spin.value()

        # Запоминаем для xlsx
        self._calc_shv = self._shv_cb.isChecked()
        self._calc_skz = self._skz_cb.isChecked() if self._calc_shv else False
        self._calc_sxx = self._sxx_cb.isChecked() if self._calc_shv else False
        self._calc_sin = self._sin_cb.isChecked()
        self._calc_equiv = calc_equiv
        self._calc_eff = calc_eff
        self._bands_list = bands
        self._base_freqs_list = base_freqs

        # Нарезаем данные по графикам × режимам
        items = []
        self._short_names: Dict[str, str] = {}
        for g_idx, g_label in checked_graphs:
            src = main_win._plot_index_to_source.get(g_idx)
            self._short_names[g_label] = src[1] if src else g_label
            data = main_win.get_plot_series_data(g_idx)
            if data is None:
                continue
            x_data, y_data = data
            x_arr = np.asarray(x_data)
            y_arr = np.asarray(y_data)
            for m_idx, m_label in checked_modes:
                _label, t0, t1 = self._modes[m_idx]
                i0 = int(np.searchsorted(x_arr, t0, side='left'))
                i1 = int(np.searchsorted(x_arr, t1, side='right'))
                x_slice = x_arr[i0:i1]
                y_slice = y_arr[i0:i1]
                if y_slice.size == 0:
                    continue
                items.append((g_label, _label, x_slice, y_slice))

        if not items:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Нет данных в выбранных режимных интервалах.")
            return

        self._results.clear()
        self._checked_graphs_labels = [lbl for _, lbl in checked_graphs]
        self._checked_modes_labels = [self._modes[mi][0] for mi, _ in checked_modes]
        self._progress.setVisible(True)
        self._progress.setValue(0)
        self._calc_btn.setEnabled(False)

        self._worker_thread = QtCore.QThread(self)
        self._worker = VibrometryBatchWorker(
            items, bands=bands, base_freqs=base_freqs,
            delta_f=delta_f, af_pct=af_pct, ae_pct=ae_pct, al_pct=al_pct,
            calc_equiv=calc_equiv, calc_eff=calc_eff,
        )
        self._worker.moveToThread(self._worker_thread)
        self._worker_thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._progress.setValue)
        self._worker.one_result.connect(self._collect_result)
        self._worker.finished.connect(self._on_finished)
        self._worker.finished.connect(self._worker_thread.quit)
        self._worker_thread.finished.connect(self._worker_thread.deleteLater)
        self._worker_thread.start()

    def _collect_result(self, graph_label: str, mode_label: str, r: FullVibrometryResult):
        self._results.append((graph_label, mode_label, r))

    def _on_finished(self):
        self._progress.setValue(100)
        self._progress.setVisible(False)
        self._calc_btn.setEnabled(True)
        if not self._results:
            QtWidgets.QMessageBox.information(self, "Готово", "Нет результатов для сохранения.")
            return
        self._export_xlsx()

    def _export_xlsx(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter as _gcl

        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Сохранить результаты", "", "Excel (*.xlsx);;Все файлы (*.*)")
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Виброметрирование"
        bold = Font(bold=True)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # Собираем структуру столбцов по одному графику:
        bands = self._bands_list
        has_skz = self._calc_skz
        has_sxx = self._calc_sxx
        n_bands = len(bands)
        n_shv_cols = 0
        if self._calc_shv:
            if has_skz:
                n_shv_cols += n_bands
            if has_sxx:
                n_shv_cols += n_bands

        base_freqs = self._base_freqs_list
        has_equiv = self._calc_equiv
        has_eff = self._calc_eff
        n_sin_per_freq = int(has_equiv) + int(has_eff)
        n_sin_cols = len(base_freqs) * n_sin_per_freq if self._calc_sin else 0

        n_cols_per_graph = n_shv_cols + n_sin_cols
        if n_cols_per_graph == 0:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "Нет столбцов для экспорта.")
            return

        graphs = self._checked_graphs_labels
        modes = self._checked_modes_labels

        # Группируем результаты
        data_map: Dict[str, Dict[str, FullVibrometryResult]] = {}
        for g, m, r in self._results:
            data_map.setdefault(g, {})[m] = r

        # A1..A4: «№ реж.» (merge 4 rows)
        ws.merge_cells(start_row=1, start_column=1, end_row=4, end_column=1)
        c0 = ws.cell(1, 1, "№ реж.")
        c0.font = bold
        c0.alignment = center

        col = 2  # начинаем со столбца 2
        for g in graphs:
            short = self._short_names.get(g, g)
            g_start = col
            g_end = col + n_cols_per_graph - 1

            # Строка 1: имя графика
            if g_end > g_start:
                ws.merge_cells(start_row=1, start_column=g_start, end_row=1, end_column=g_end)
            c = ws.cell(1, g_start, short)
            c.font = bold
            c.alignment = center

            inner_col = g_start

            if self._calc_shv and n_shv_cols > 0:
                shv_start = inner_col
                shv_end = inner_col + n_shv_cols - 1
                # Строка 2: «С.Ш.В.»
                if shv_end > shv_start:
                    ws.merge_cells(start_row=2, start_column=shv_start, end_row=2, end_column=shv_end)
                c = ws.cell(2, shv_start, "С.Ш.В.")
                c.font = bold
                c.alignment = center

                # Строка 3: «СКЗ по полосам» / «Sxx по полосам»
                if has_skz:
                    skz_start = inner_col
                    skz_end = inner_col + n_bands - 1
                    if skz_end > skz_start:
                        ws.merge_cells(start_row=3, start_column=skz_start, end_row=3, end_column=skz_end)
                    c = ws.cell(3, skz_start, "СКЗ по полосам")
                    c.font = bold
                    c.alignment = center
                    # Строка 4: диапазоны
                    for bi, (f1, f2) in enumerate(bands):
                        lbl = f"{f1:g}...{f2:g}"
                        c = ws.cell(4, inner_col + bi, lbl)
                        c.font = bold
                        c.alignment = center
                    inner_col += n_bands

                if has_sxx:
                    sxx_start = inner_col
                    sxx_end = inner_col + n_bands - 1
                    if sxx_end > sxx_start:
                        ws.merge_cells(start_row=3, start_column=sxx_start, end_row=3, end_column=sxx_end)
                    c = ws.cell(3, sxx_start, "Sxx по полосам")
                    c.font = bold
                    c.alignment = center
                    for bi, (f1, f2) in enumerate(bands):
                        lbl = f"{f1:g}...{f2:g}"
                        c = ws.cell(4, inner_col + bi, lbl)
                        c.font = bold
                        c.alignment = center
                    inner_col += n_bands

            if self._calc_sin and n_sin_cols > 0:
                sin_start = inner_col
                sin_end = inner_col + n_sin_cols - 1
                # Строка 2: «Синусоидальная вибрация»
                if sin_end > sin_start:
                    ws.merge_cells(start_row=2, start_column=sin_start, end_row=2, end_column=sin_end)
                c = ws.cell(2, sin_start, "Синусоидальная вибрация")
                c.font = bold
                c.alignment = center

                # Строка 3: по каждой базовой частоте
                for fi, f0 in enumerate(base_freqs):
                    freq_start = inner_col
                    freq_end = inner_col + n_sin_per_freq - 1
                    freq_label = f"{f0:g} Гц"
                    if freq_end > freq_start:
                        ws.merge_cells(start_row=3, start_column=freq_start, end_row=3, end_column=freq_end)
                    c = ws.cell(3, freq_start, freq_label)
                    c.font = bold
                    c.alignment = center

                    # Строка 4: Экв. / Эфф.
                    sub_col = inner_col
                    if has_equiv:
                        c = ws.cell(4, sub_col, "Экв.")
                        c.font = bold
                        c.alignment = center
                        sub_col += 1
                    if has_eff:
                        c = ws.cell(4, sub_col, "Эфф.")
                        c.font = bold
                        c.alignment = center
                        sub_col += 1
                    inner_col += n_sin_per_freq

            col += n_cols_per_graph

        DATA_START_ROW = 5
        for ri, mode_label in enumerate(modes):
            row = DATA_START_ROW + ri
            mc = ws.cell(row, 1, mode_label)
            mc.font = bold
            mc.alignment = center
            col = 2
            for g in graphs:
                r = data_map.get(g, {}).get(mode_label)
                inner_col = col

                # С.Ш.В.
                if self._calc_shv:
                    if has_skz:
                        for bi in range(n_bands):
                            val = ""
                            if r and bi < len(r.bands):
                                v = r.bands[bi].rms
                                val = v if np.isfinite(v) else ""
                            c = ws.cell(row, inner_col + bi, val)
                            c.alignment = center
                        inner_col += n_bands
                    if has_sxx:
                        for bi in range(n_bands):
                            val = ""
                            if r and bi < len(r.bands):
                                v = r.bands[bi].sxx
                                val = v if np.isfinite(v) else ""
                            c = ws.cell(row, inner_col + bi, val)
                            c.alignment = center
                        inner_col += n_bands

                # Синусоидальная
                if self._calc_sin:
                    for fi in range(len(base_freqs)):
                        sr = r.sinusoidal[fi] if r and fi < len(r.sinusoidal) else None
                        if has_equiv:
                            val = ""
                            if sr and sr.equiv_amplitude is not None and np.isfinite(sr.equiv_amplitude):
                                val = sr.equiv_amplitude
                            c = ws.cell(row, inner_col, val)
                            c.alignment = center
                            inner_col += 1
                        if has_eff:
                            val = ""
                            if sr and sr.eff_amplitude is not None and np.isfinite(sr.eff_amplitude):
                                val = sr.eff_amplitude
                            c = ws.cell(row, inner_col, val)
                            c.alignment = center
                            inner_col += 1

                col += n_cols_per_graph

        # Автоширина
        for ci in range(1, ws.max_column + 1):
            max_len = 0
            for ri in range(1, ws.max_row + 1):
                val = ws.cell(ri, ci).value
                if val is not None:
                    s = f"{val:.6g}" if isinstance(val, float) else str(val)
                    max_len = max(max_len, len(s))
            ws.column_dimensions[_gcl(ci)].width = max(8, max_len + 3)

        try:
            wb.save(path)
            QtWidgets.QMessageBox.information(self, "Готово", f"Результаты сохранены:\n{path}")
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Ошибка сохранения", str(e))


class CalculationDialog(QtWidgets.QDialog):
    def __init__(self, title: str, items: list[dict], parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(420, 380)

        # items: список словарей с полями как минимум "label" и "plot_idx".
        self._items = items
        self._worker_thread = None
        self._worker = None

        layout = QtWidgets.QVBoxLayout(self)
        self._stack = QtWidgets.QStackedWidget()
        layout.addWidget(self._stack)

        self._build_graph_page()
        self._page_mode = self._build_mode_page()
        self._page_format = self._build_format_page()
        self._page_run = self._build_run_page()

        self._stack.addWidget(self._page_graph)
        self._stack.addWidget(self._page_mode)
        self._stack.addWidget(self._page_format)
        self._stack.addWidget(self._page_run)

        btns = QtWidgets.QHBoxLayout()
        self._btn_back = QtWidgets.QPushButton("Назад")
        self._btn_next = QtWidgets.QPushButton("Продолжить")
        self._btn_close = QtWidgets.QPushButton("Закрыть")
        btns.addWidget(self._btn_back)
        btns.addWidget(self._btn_next)
        btns.addStretch(1)
        btns.addWidget(self._btn_close)
        layout.addLayout(btns)

        self._btn_back.clicked.connect(self._prev_step)
        self._btn_next.clicked.connect(self._next_step)
        self._btn_close.clicked.connect(self.reject)

        self._update_nav()

    def _build_graph_page(self) -> QtWidgets.QWidget:
        self._page_graph = QtWidgets.QWidget()
        lay = QtWidgets.QVBoxLayout(self._page_graph)
        lay.addWidget(QtWidgets.QLabel("Шаг 1. Выбор графика"))
        self._graph_combo = QtWidgets.QComboBox()
        if self._items:
            for item in self._items:
                label = item.get("label", "")
                self._graph_combo.addItem(label, item)
        else:
            self._graph_combo.addItem("Нет активных графиков")
            self._graph_combo.setEnabled(False)
        lay.addWidget(self._graph_combo)
        lay.addStretch(1)
        return self._page_graph

    def _build_mode_page(self) -> QtWidgets.QWidget:
        page = QtWidgets.QWidget()
        lay = QtWidgets.QVBoxLayout(page)
        lay.addWidget(QtWidgets.QLabel("Шаг 2. Режим и параметры"))
        self._mode_combo = QtWidgets.QComboBox()
        self._mode_combo.addItem("Основной режим")
        lay.addWidget(self._mode_combo)
        lay.addWidget(QtWidgets.QLabel("Параметры отсутствуют"))
        lay.addStretch(1)
        return page

    def _build_format_page(self) -> QtWidgets.QWidget:
        page = QtWidgets.QWidget()
        lay = QtWidgets.QVBoxLayout(page)
        lay.addWidget(QtWidgets.QLabel("Шаг 3. Формат результата"))
        self._format_combo = QtWidgets.QComboBox()
        self._format_combo.addItems([
            "Добавить как новый график",
            "Сохранить в файл",
        ])
        lay.addWidget(self._format_combo)
        lay.addStretch(1)
        return page

    def _build_run_page(self) -> QtWidgets.QWidget:
        page = QtWidgets.QWidget()
        lay = QtWidgets.QVBoxLayout(page)
        lay.addWidget(QtWidgets.QLabel("Шаг 4. Расчет"))
        self._progress = QtWidgets.QProgressBar()
        self._progress.setRange(0, 100)
        self._progress.setValue(0)
        self._status = QtWidgets.QLabel("Готово к расчету")
        self._calc_btn = QtWidgets.QPushButton("Рассчитать")
        self._calc_btn.clicked.connect(self._start_calculation)
        self._results_text = QtWidgets.QTextEdit()
        self._results_text.setReadOnly(True)
        self._results_text.setPlaceholderText("Результаты появятся после расчета.")
        self._results_text.setMinimumHeight(120)
        lay.addWidget(self._progress)
        lay.addWidget(self._status)
        lay.addWidget(self._calc_btn)
        lay.addWidget(QtWidgets.QLabel("Результаты:"))
        lay.addWidget(self._results_text)
        lay.addStretch(1)
        return page

    def _update_nav(self):
        idx = self._stack.currentIndex()
        self._btn_back.setEnabled(idx > 0)
        self._btn_next.setEnabled(idx < self._stack.count() - 1)
        if not self._items:
            self._btn_next.setEnabled(False)

    def _next_step(self):
        idx = self._stack.currentIndex()
        if idx < self._stack.count() - 1:
            self._stack.setCurrentIndex(idx + 1)
        self._update_nav()

    def _prev_step(self):
        idx = self._stack.currentIndex()
        if idx > 0:
            self._stack.setCurrentIndex(idx - 1)
        self._update_nav()

    def _start_calculation(self):
        if self._worker_thread is not None:
            return
        self._start_generic()

    def _start_generic(self):
        self._progress.setValue(0)
        self._status.setText("Выполняется расчет...")
        self._calc_btn.setEnabled(False)
        self._btn_back.setEnabled(False)
        self._btn_next.setEnabled(False)

        self._worker_thread = QtCore.QThread(self)
        self._worker = CalculationWorker()
        self._worker.moveToThread(self._worker_thread)
        self._worker_thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._progress.setValue)
        self._worker.finished.connect(self._on_calc_finished)
        self._worker.finished.connect(self._worker_thread.quit)
        self._worker_thread.finished.connect(self._worker_thread.deleteLater)
        self._worker_thread.start()

    def _on_calc_finished(self, message: str):
        self._status.setText(message)
        self._btn_back.setEnabled(True)
        self._btn_next.setEnabled(True)
        self._btn_close.setEnabled(True)
        self._worker_thread = None
        self._worker = None
